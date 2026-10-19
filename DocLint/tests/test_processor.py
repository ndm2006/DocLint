import base64
from io import BytesIO
from pathlib import Path
import unittest
from unittest.mock import patch
import warnings
import zipfile
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches
from lxml import etree
from core.autofix import DocFixer
from core.processor import DocumentError, process_document
from core.xmlutils import NS, q, parse_xml, content_signature

BASE = Path(__file__).resolve().parents[1]
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')

def save(document):
    out = BytesIO(); document.save(out)
    return out.getvalue()

def replace_part(data, name, replacement):
    out = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as src, zipfile.ZipFile(out, 'w') as dst:
        for info in src.infolist():
            dst.writestr(info.filename, replacement if info.filename == name else src.read(info.filename))
    return out.getvalue()

def document_xml(data):
    with zipfile.ZipFile(BytesIO(data)) as z:
        return parse_xml(z.read('word/document.xml'))

class ProcessorTests(unittest.TestCase):
    def setUp(self):
        self.doc = Document()
        self.doc.add_paragraph('Đây là phần thân bài có dấu tiếng Việt, đường dẫn và dữ liệu gốc.')
        self.data = save(self.doc)

    def test_actual_formatting_and_idempotence(self):
        result = process_document(self.data)
        xml = document_xml(result.data)
        self.assertTrue(result.report['integrity']['verified'])
        self.assertGreater(result.report['summary']['changed_properties'], 0)
        for path, key, expected in [('.//w:pgSz','w','11906'),('.//w:pgMar','top','1134'),('.//w:rFonts','ascii','Times New Roman'),('.//w:spacing','line','360'),('.//w:sz','val','26')]:
            self.assertEqual(xml.find(path, NS).get(q(key)), expected)
        self.assertFalse(process_document(result.data).report['changes'])
        self.assertEqual(Document(BytesIO(result.data)).paragraphs[0].text, self.doc.paragraphs[0].text)

    def test_every_character_blank_paragraph_and_table_is_preserved(self):
        self.doc.add_paragraph('  Khoảng trắng ở đầu và cuối  ')
        self.doc.add_paragraph('\tHà Nội\tngày 12 tháng 10')
        blank = self.doc.add_paragraph()
        bookmark = OxmlElement('w:bookmarkStart'); bookmark.set(q('id'), '7'); bookmark.set(q('name'), 'keep')
        blank._p.append(bookmark)
        self.doc.add_paragraph()
        table = self.doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Ô gộp'; table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(1, 0).add_table(rows=1, cols=1)
        equation = OxmlElement('m:oMath')
        run = OxmlElement('m:r'); text = OxmlElement('m:t'); text.text = 'E = mc²'
        run.append(text); equation.append(run); self.doc.add_paragraph()._p.append(equation)
        self.doc.add_picture(BytesIO(PNG), width=Inches(.5))
        self.doc.sections[0].header.paragraphs[0].text = 'Đầu trang được giữ nguyên'
        original = save(self.doc); result = process_document(original)
        before, after = document_xml(original), document_xml(result.data)
        self.assertEqual(before.xpath('//w:t/text()', namespaces=NS), after.xpath('//w:t/text()', namespaces=NS))
        self.assertEqual(len(before.xpath('//w:p', namespaces=NS)), len(after.xpath('//w:p', namespaces=NS)))
        for path in ('//w:tbl', '//m:oMath', '//w:drawing', '//w:tab', '//w:bookmarkStart'):
            self.assertEqual([etree.tostring(n) for n in before.xpath(path, namespaces=NS)], [etree.tostring(n) for n in after.xpath(path, namespaces=NS)])
        self.assertEqual(content_signature(before), content_signature(after))
        with zipfile.ZipFile(BytesIO(original)) as src, zipfile.ZipFile(BytesIO(result.data)) as dst:
            self.assertEqual(src.namelist(), dst.namelist())
            for name in src.namelist():
                if name not in {'word/document.xml', 'word/styles.xml'}:
                    self.assertEqual(src.read(name), dst.read(name), name)
        self.assertEqual(result.report['summary']['equations'], 1)
        self.assertEqual(result.report['summary']['images'], 1)
        self.assertEqual(result.report['summary']['tables'], 2)

    def test_risky_heading_and_existing_heading_are_not_reassigned(self):
        fake = self.doc.add_paragraph('Chương 1: tổng quan dự án'); fake.runs[0].bold = True
        heading = self.doc.add_heading('Mục đã có cấu trúc', level=2)
        before = [etree.tostring(fake._p), etree.tostring(heading._p)]
        result = process_document(save(self.doc))
        paragraphs = document_xml(result.data).findall('w:body/w:p', NS)
        self.assertEqual(before, [etree.tostring(p) for p in paragraphs[-2:]])
        self.assertIn('headings', [w['code'] for w in result.report['warnings']])

    def test_landscape_columns_fields_and_symbols_stay_unchanged(self):
        section = self.doc.sections[0]; section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
        section._sectPr.find(q('cols')).set(q('num'), '2')
        field = self.doc.add_paragraph()
        field_char = OxmlElement('w:fldSimple'); field_char.set(q('instr'), ' PAGE '); field._p.append(field_char)
        symbol = self.doc.add_paragraph('αβγ'); symbol.runs[0].font.name = 'Symbol'
        source = save(self.doc); result = process_document(source)
        a, b = document_xml(source), document_xml(result.data)
        for path in ('//w:sectPr', '//w:fldSimple'):
            self.assertEqual(etree.tostring(a.xpath(path, namespaces=NS)[0]), etree.tostring(b.xpath(path, namespaces=NS)[0]))
        self.assertEqual(b.xpath('//w:rFonts/@w:ascii', namespaces=NS)[-1], 'Symbol')
        self.assertTrue({'complex_sections', 'symbol_fonts'}.issubset({w['code'] for w in result.report['warnings']}))

    def test_existing_signature_alignment_is_preserved(self):
        p = self.doc.add_paragraph('Sinh viên thực hiện'); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        result = process_document(save(self.doc))
        self.assertEqual(document_xml(result.data).xpath('//w:pPr/w:jc/@w:val', namespaces=NS)[-1], 'right')

    def test_integrity_gate_rejects_accidental_content_mutation(self):
        original_fix = DocFixer.fix_paragraphs
        def bad_fix(fixer, tree):
            original_fix(fixer, tree); tree.find('.//w:t', NS).text = 'Nội dung bị thay đổi'
            return tree
        with patch.object(DocFixer, 'fix_paragraphs', bad_fix), self.assertRaises(DocumentError):
            process_document(self.data)

    def test_invalid_zip_wrong_types_entities_and_duplicates_rejected(self):
        with self.assertRaises(DocumentError): process_document(b'not a docx')
        for replacement in (b'<!DOCTYPE document [<!ENTITY e SYSTEM "file:///etc/passwd">]><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>&e;</w:t></w:r></w:p></w:body></w:document>', b'<broken'):
            with self.assertRaises(DocumentError): process_document(replace_part(self.data, 'word/document.xml', replacement))
        with self.assertRaises(DocumentError): process_document(replace_part(self.data, '[Content_Types].xml', b'<Types/>'))
        out = BytesIO(self.data)
        with zipfile.ZipFile(out, 'a') as archive, warnings.catch_warnings():
            warnings.simplefilter('ignore', UserWarning); archive.writestr('word/document.xml', b'fake')
        with self.assertRaises(DocumentError): process_document(out.getvalue())

    def test_size_limits_and_signed_packages(self):
        for name in ('MAX_EXPANDED', 'MAX_UPLOAD', 'MAX_XML'):
            with patch('core.processor.' + name, 1), self.assertRaises(DocumentError): process_document(self.data)
        for name in ('word/vbaProject.bin', '_xmlsignatures/sig1.xml', '../escape.xml'):
            out = BytesIO(self.data)
            with zipfile.ZipFile(out, 'a') as z: z.writestr(name, b'not supported')
            with self.assertRaises(DocumentError): process_document(out.getvalue())

    def test_both_profiles_and_unknown_profile(self):
        result = process_document(self.data, 'report'); tree = document_xml(result.data)
        self.assertEqual(tree.find('.//w:rFonts', NS).get(q('ascii')), 'Calibri')
        self.assertEqual(tree.find('.//w:spacing', NS).get(q('line')), '288')
        self.assertFalse(process_document(result.data, 'report').report['changes'])
        with self.assertRaises(ValueError): process_document(self.data, 'ieee')

    def test_repository_fixture_preserves_all_text_exactly(self):
        data = (BASE / 'test_input.docx').read_bytes(); before = document_xml(data)
        result = process_document(data)
        self.assertEqual(before.xpath('//w:t/text()', namespaces=NS), document_xml(result.data).xpath('//w:t/text()', namespaces=NS))
        self.assertEqual(content_signature(before), content_signature(document_xml(result.data)))
