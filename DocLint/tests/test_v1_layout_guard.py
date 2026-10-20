"""Preserve fixed page units; format only independent flowing text."""

from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from fastapi.testclient import TestClient

from api import create_app
from core.processor import DocumentError
from core.xmlutils import NS, q
from v1.analysis import Analysis, text_of
from v1.editor import Editor
from v1.package import Package, canonical
from v1.pipeline import format_document
from v1.renderer import find_office
from test_v1_jobs import wait_for, unavailable
from v1_fixtures import fixture, save


def frame(paragraph, y=1800):
    node = OxmlElement("w:framePr")
    for key, value in {
        "w": 7000,
        "x": 1440,
        "y": y,
        "hAnchor": "page",
        "vAnchor": "page",
    }.items():
        node.set(q(key), str(value))
    paragraph._p.get_or_add_pPr().append(node)
    return node


def positioned_doc():
    d = Document()
    frame(d.add_paragraph("UNIVERSITY OF EXAMPLE"))
    frame(d.add_paragraph("ESSAY"), 3000)
    frame(d.add_paragraph("Subject: An example topic"), 4200)
    d.add_page_break()
    frame(d.add_paragraph("A fixed content line"))
    frame(d.add_paragraph("Another fixed content line"), 2300)
    return d


def mixed_doc():
    d = Document()
    frame(d.add_paragraph("UNIVERSITY OF EXAMPLE"))
    frame(d.add_paragraph("ESSAY"), 3000)
    d.add_page_break()
    d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
    d.add_paragraph("A normal paragraph that can be formatted independently.")
    d.add_page_break()
    frame(d.add_paragraph("An appendix on its own fixed page"))
    return d


class LayoutGuardTests(unittest.TestCase):
    def test_all_positioned_units_preserve_the_original_file_byte_for_byte(self):
        data = save(positioned_doc())
        a = Analysis(Package(data))
        self.assertEqual(a.layout_policy["kind"], "fixed")
        self.assertEqual(a.layout_policy["frames"], 5)
        self.assertEqual(len(a.layout_policy["protected_regions"]), 2)
        self.assertEqual(a.cover_end, "p00004")
        self.assertEqual(a.public()["review_plan"]["question_count"], 0)
        editor = Editor(
            data,
            {
                "mode": "layout",
                "cover": "academic",
                "toc": "insert",
                "page_numbers": "academic",
                "object_lists": True,
                "caption_fields": True,
            },
        ).run()
        self.assertEqual(editor.changes, [])
        self.assertEqual(editor.serialize(), data)
        self.assertTrue(editor.ledger.verify_bytes(editor.serialize()))
        self.assertEqual(editor.report()["operation_mode"], "preserve_only")

    def test_mixed_document_preserves_fixed_units_and_shared_section(self):
        data = save(mixed_doc())
        source = Analysis(Package(data))
        editor = Editor(
            data,
            {
                "mode": "layout",
                "cover": "academic",
                "toc": "insert",
                "page_numbers": "academic",
                "object_lists": True,
                "caption_fields": True,
                "convert_manual_lists": True,
                "signature_position": "end_of_body",
            },
        ).run()
        self.assertEqual(editor.report()["operation_mode"], "selective")
        self.assertTrue(editor.changes)
        for bid in source.layout_policy["protected_block_ids"]:
            self.assertEqual(
                canonical(editor.analysis.elements[bid]),
                canonical(source.elements[bid]),
                bid,
            )
        self.assertEqual(
            canonical(editor.analysis.section_nodes[0][1]),
            canonical(source.section_nodes[0][1]),
        )
        self.assertNotEqual(
            canonical(editor.analysis.elements["p00005"]),
            canonical(source.elements["p00005"]),
        )
        self.assertEqual(
            editor.analysis.elements["p00005"]
            .find("w:pPr/w:spacing", NS)
            .get(q("line")),
            "360",
        )
        self.assertFalse(editor.package.document.findall(".//" + q("sdt")))
        self.assertTrue(editor.ledger.verify_bytes(editor.serialize()))
        self.assertEqual(
            text_of(editor.package.document), text_of(source.package.document)
        )

    def test_frame_inherited_from_style_is_also_protected(self):
        d = Document()
        style = d.styles.add_style("PositionedText", WD_STYLE_TYPE.PARAGRAPH)
        props = OxmlElement("w:pPr")
        props.append(deepcopy(frame(d.add_paragraph("Styled fixed content"))))
        style.element.append(props)
        p = d.paragraphs[0]
        p._p.get_or_add_pPr().remove(p._p.pPr.find(q("framePr")))
        p.style = style
        a = Analysis(Package(save(d)))
        self.assertTrue(a.layout_policy["preserve_only"])

    def test_rendered_break_cache_cannot_separate_a_safe_unit(self):
        d = Document()
        frame(d.add_paragraph("Fixed text"))
        p = d.add_paragraph("Text sharing that page")
        p.runs[0]._r.insert(0, OxmlElement("w:lastRenderedPageBreak"))
        a = Analysis(Package(save(d)))
        self.assertEqual(a.layout_policy["protected_block_ids"], ["p00001", "p00002"])
        self.assertTrue(a.layout_policy["preserve_only"])

    def test_page_break_before_delimits_the_next_independent_unit(self):
        d = Document()
        frame(d.add_paragraph("Fixed text"))
        p = d.add_paragraph("Flowing text on a new page")
        p.paragraph_format.page_break_before = True
        a = Analysis(Package(save(d)))
        self.assertEqual(a.layout_policy["protected_block_ids"], ["p00001"])
        self.assertEqual(a.layout_policy["editable_block_ids"], ["p00002"])

    def test_continuous_section_cannot_split_a_protected_page(self):
        d = Document()
        frame(d.add_paragraph("A fixed region"))
        d.add_section(WD_SECTION_START.CONTINUOUS)
        d.add_paragraph("Ordinary text sharing that page")
        a = Analysis(Package(save(d)))
        self.assertTrue(a.layout_policy["preserve_only"])
        self.assertEqual(a.layout_policy["protected_sections"], [0, 1])

    def test_next_page_section_can_delimit_an_independent_flow_region(self):
        d = Document()
        frame(d.add_paragraph("A fixed region"))
        d.add_section(WD_SECTION_START.NEW_PAGE)
        d.add_paragraph("Ordinary text on a separate page")
        a = Analysis(Package(save(d)))
        self.assertFalse(a.layout_policy["preserve_only"])
        self.assertEqual(a.layout_policy["editable_block_ids"], ["p00003"])

    def test_inline_textbox_protects_surrounding_page_unit(self):
        d = Document()
        p = d.add_paragraph("Text sharing the page with a text box")
        pict = OxmlElement("w:pict")
        box = OxmlElement("w:txbxContent")
        box_p = OxmlElement("w:p")
        box.append(box_p)
        pict.append(box)
        p.add_run()._r.append(pict)
        a = Analysis(Package(save(d)))
        self.assertEqual(a.layout_policy["textboxes"], 1)
        self.assertTrue(a.layout_policy["preserve_only"])

    def test_floating_table_is_not_normalized_as_a_flowing_data_table(self):
        d = Document()
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Fixed table"
        table._tbl.tblPr.append(OxmlElement("w:tblpPr"))
        a = Analysis(Package(save(d)))
        self.assertEqual(a.layout_policy["floating_tables"], 1)
        self.assertTrue(a.layout_policy["preserve_only"])

    def test_break_inside_textbox_is_not_a_main_story_boundary(self):
        d = Document()
        p = d.add_paragraph()
        pict = OxmlElement("w:pict")
        box = OxmlElement("w:txbxContent")
        box_p = OxmlElement("w:p")
        run = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(q("type"), "page")
        run.append(br)
        box_p.append(run)
        box.append(box_p)
        pict.append(box)
        p.add_run()._r.append(pict)
        d.add_paragraph("This paragraph still shares the text box's page.")
        a = Analysis(Package(save(d)))
        self.assertEqual(a.layout_policy["protected_block_ids"], ["p00001", "p00002"])

    def test_paragraph_attached_object_does_not_disable_headings_in_surrounding_text(
        self,
    ):
        d = Document()
        d.add_heading("NỘI DUNG", 1)
        p = d.add_paragraph()
        drawing = OxmlElement("w:drawing")
        anchor = OxmlElement("wp:anchor")
        vertical = OxmlElement("wp:positionV")
        vertical.set("relativeFrom", "paragraph")
        anchor.append(vertical)
        box = OxmlElement("w:txbxContent")
        box.append(OxmlElement("w:p"))
        anchor.append(box)
        drawing.append(anchor)
        p.add_run()._r.append(drawing)
        d.add_paragraph("Ordinary text outside the graphic.")
        data = save(d)
        a = Analysis(Package(data))
        self.assertEqual(a.layout_policy["kind"], "flow")
        self.assertEqual(a.layout_policy["isolated_object_ids"], ["p00002"])
        e = Editor(data).run()
        self.assertEqual(
            canonical(e.analysis.elements["p00002"]), canonical(a.elements["p00002"])
        )
        self.assertNotEqual(
            canonical(e.analysis.elements["p00003"]), canonical(a.elements["p00003"])
        )
        self.assertTrue(e.ledger.verify_bytes(e.serialize()))

    def test_drop_cap_does_not_protect_an_otherwise_ordinary_report(self):
        d = Document()
        p = d.add_paragraph("An ordinary paragraph with a drop cap")
        f = OxmlElement("w:framePr")
        f.set(q("dropCap"), "drop")
        p._p.get_or_add_pPr().append(f)
        a = Analysis(Package(save(d)))
        self.assertEqual(a.layout_policy["kind"], "flow")

    def test_unchanged_candidate_and_body_conversion_are_distinct(self):
        d = Document()
        p = d.add_paragraph("AN IMPORTANT TOPIC NAME")
        p.runs[0].bold = True
        p.paragraph_format.alignment = 1
        d.add_paragraph("A paragraph explaining that topic.")
        data = save(d)
        original = Analysis(Package(data))
        keep = Editor(data, {"unchanged_candidates": ["p00001"]}).run()
        body = Editor(data, {"preserved_candidates": ["p00001"]}).run()
        heading = Editor(data, {"heading_levels": {"p00001": 2}}).run()
        self.assertEqual(
            canonical(keep.analysis.elements["p00001"]),
            canonical(original.elements["p00001"]),
        )
        self.assertNotEqual(
            canonical(body.analysis.elements["p00001"]),
            canonical(original.elements["p00001"]),
        )
        self.assertIsNone(
            body.analysis.elements["p00001"].find("w:pPr/w:outlineLvl", NS)
        )
        self.assertEqual(
            heading.analysis.elements["p00001"]
            .find("w:pPr/w:outlineLvl", NS)
            .get(q("val")),
            "1",
        )
        self.assertTrue(keep.ledger.verify_bytes(keep.serialize()))
        with self.assertRaises(DocumentError):
            Editor(
                data,
                {"unchanged_candidates": ["p00001"], "heading_levels": {"p00001": 1}},
            )
        with self.assertRaises(DocumentError):
            Editor(
                data,
                {
                    "unchanged_candidates": ["p00001"],
                    "preserved_candidates": ["p00001"],
                },
            )

    def test_invalid_or_certain_heading_cannot_be_an_unchanged_candidate(self):
        d = Document()
        d.add_heading("A native heading", 1)
        data = save(d)
        for bid in ("p00001", "missing"):
            with self.assertRaises(DocumentError):
                Editor(data, {"unchanged_candidates": [bid]})

    def test_english_flow_cover_is_recognized_before_the_first_page_break(self):
        d = Document()
        d.add_paragraph("UNIVERSITY OF EXAMPLE")
        d.add_paragraph("FINAL ESSAY")
        d.add_paragraph("Subject: An example topic")
        d.add_paragraph("Instructor: Example Name")
        d.add_page_break()
        d.add_heading("Introduction", 1)
        a = Analysis(Package(save(d)))
        self.assertEqual(a.cover_end, "p00005")
        self.assertEqual(a.layout_policy["kind"], "flow")

    def test_both_cover_templates_center_every_safe_line_including_heading_styles(self):
        d = Document()
        d.add_paragraph("TRƯỜNG ĐẠI HỌC MINH HỌA")
        d.add_paragraph("KHOA THỬ NGHIỆM")
        d.add_paragraph()
        d.add_heading("BÁO CÁO MÔN HỌC", 1)
        d.add_heading("TÊN CHỦ ĐỀ ĐƯỢC GIỮ NGUYÊN", 2)
        d.add_paragraph("Giảng viên: Tên minh họa")
        d.add_paragraph("Nhóm thực hiện: Nhóm A")
        d.add_page_break()
        d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
        data = save(d)
        a = Analysis(Package(data))
        self.assertEqual(a.cover_end, "p00008")
        self.assertTrue(a.cover_explicit)
        for template in ("academic", "minimal"):
            e = Editor(data, {"cover": template, "mode": "layout"}).run()
            for bid in ("p00001", "p00002", "p00004", "p00005", "p00006", "p00007"):
                p = e.analysis.elements[bid]
                self.assertEqual(
                    p.find("w:pPr/w:jc", NS).get(q("val")), "center", (template, bid)
                )
                self.assertEqual(text_of(p), text_of(a.elements[bid]))
            self.assertEqual(
                e.analysis.elements["p00009"].find("w:pPr/w:jc", NS).get(q("val")),
                "left",
            )
            self.assertTrue(e.ledger.verify_bytes(e.serialize()))

    def test_intro_prose_is_not_swallowed_by_a_later_body_page_break(self):
        d = Document()
        d.add_heading("BÁO CÁO MÔN HỌC", 1)
        d.add_paragraph("Giảng viên: Tên minh họa")
        d.add_paragraph("Đây là phần mở đầu thân bài cần được giữ bên ngoài trang bìa.")
        d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
        d.add_paragraph("Phần trình bày nội dung.")
        d.add_page_break()
        a = Analysis(Package(save(d)))
        self.assertEqual(a.cover_end, "p00002")

    def test_page_break_before_body_is_an_explicit_cover_boundary(self):
        d = Document()
        d.add_paragraph("UNIVERSITY OF EXAMPLE")
        d.add_paragraph("FINAL ESSAY")
        p = d.add_heading("Introduction", 1)
        p.paragraph_format.page_break_before = True
        a = Analysis(Package(save(d)))
        self.assertEqual(a.cover_end, "p00002")
        self.assertTrue(a.cover_explicit)

    def test_english_cover_title_receives_cover_typography(self):
        d = Document()
        d.add_paragraph("UNIVERSITY OF EXAMPLE")
        d.add_paragraph("FINAL ESSAY")
        d.add_paragraph("Subject: An example topic")
        d.add_page_break()
        d.add_heading("Introduction", 1)
        e = Editor(save(d), {"cover": "academic"}).run()
        title = e.analysis.elements["p00002"]
        self.assertEqual(title.find("w:pPr/w:jc", NS).get(q("val")), "center")
        self.assertEqual(title.find(".//" + q("sz")).get(q("val")), "40")

    def test_case_00_cover_is_centered_with_both_templates_and_landscape_is_preserved(
        self,
    ):
        data = fixture(0)
        source = Analysis(Package(data))
        self.assertEqual(source.cover_end, "p00004")
        self.assertTrue(source.sections[-1]["complex"])
        for template in ("academic", "minimal"):
            e = Editor(data, {"cover": template, "mode": "layout"}).run()
            for bid in ("p00001", "p00002", "p00003"):
                self.assertEqual(
                    e.analysis.elements[bid].find("w:pPr/w:jc", NS).get(q("val")),
                    "center",
                    (template, bid),
                )
            self.assertEqual(
                canonical(e.analysis.section_nodes[-1][1]),
                canonical(source.section_nodes[-1][1]),
            )
            self.assertTrue(e.ledger.verify_bytes(e.serialize()))

    def test_user_can_reject_a_cover_proposal_without_it_being_restored(self):
        d = Document()
        d.add_heading("BÁO CÁO MÔN HỌC", 1)
        d.add_heading("TỔNG QUAN", 1)
        data = save(d)
        self.assertEqual(Analysis(Package(data)).cover_end, "p00001")
        e = Editor(data, {"cover_end": "", "cover": "keep"}).run()
        self.assertIsNone(e.cover_end)
        self.assertNotIn("p00001", e.locked)
        self.assertIn("p00001", [b["id"] for b in e.heading_blocks])
        self.assertTrue(e.ledger.verify_bytes(e.serialize()))

    @unittest.skipUnless(find_office(), "LibreOffice is required")
    def test_preserve_only_output_is_original_bytes_and_is_not_claimed_formatted(self):
        data = save(positioned_doc())
        with tempfile.TemporaryDirectory() as directory:
            output, report = format_document(
                data, {"toc": "insert", "page_numbers": "academic"}, directory
            )
            self.assertEqual(output, data)
            self.assertEqual(Path(directory, "result.docx").read_bytes(), data)
            self.assertEqual(report["operation_mode"], "preserve_only")
            self.assertEqual(report["summary"]["changed_items"], 0)
            self.assertEqual(report["status"], "partial")
            self.assertTrue(report["integrity"]["verified"])
            self.assertIn("layout_preservation", [c["code"] for c in report["checks"]])

    def test_api_returns_policy_and_true_keep_choice_without_changing_candidate(self):
        d = Document()
        d.add_paragraph("A TOPIC NAME").runs[0].bold = True
        d.add_paragraph("An ordinary explanation paragraph.")
        data = save(d)
        with (
            tempfile.TemporaryDirectory() as directory,
            patch("v1.jobs.render_pdf", unavailable),
            patch("v1.pipeline.render_pdf", unavailable),
            TestClient(create_app(directory)) as client,
        ):
            job = client.post(
                "/api/v1/jobs", files={"file": ("example.docx", data)}
            ).json()
            url = job["urls"]["status"]
            meta = wait_for(client, url)
            self.assertEqual(meta["analysis"]["layout_policy"]["kind"], "flow")
            self.assertEqual(
                client.post(
                    url + "/format", json={"unchanged_candidates": ["p00001"]}
                ).status_code,
                202,
            )
            meta = wait_for(client, url)
            result = client.get(url + "/download").content
            self.assertEqual(meta["options"]["unchanged_candidates"], ["p00001"])
            self.assertEqual(
                canonical(Analysis(Package(result)).elements["p00001"]),
                canonical(Analysis(Package(data)).elements["p00001"]),
            )
