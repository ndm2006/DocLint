"""Synthetic documents only; never commit private uploaded reports."""

import base64
from io import BytesIO
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.shared import Inches
from core.xmlutils import q

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII="
)


def save(doc):
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def fixture(case=0):
    doc = Document()
    if case % 3 == 0:
        doc.add_paragraph("TRƯỜNG ĐẠI HỌC THỬ NGHIỆM").alignment = 1
        doc.add_paragraph("BÁO CÁO MÔN HỌC").alignment = 1
        doc.add_paragraph("Đề tài: Bảo toàn dữ liệu Word").alignment = 1
        doc.add_page_break()
    p = doc.add_paragraph("CHƯƠNG 1. TỔNG QUAN")
    p.runs[0].bold = True
    p = doc.add_paragraph("1.1. Mục tiêu và phạm vi")
    p.runs[0].bold = True
    p = doc.add_paragraph(
        ("   " if case % 2 else "")
        + "Đây là nội dung gốc cần giữ nguyên: tiếng Việt, 42, 3.14 và dấu câu. "
    )
    p.add_run("Phần nhấn mạnh vẫn được giữ.").italic = True
    if case % 5 == 0:
        doc.add_paragraph("\tKhông được xóa Tab bên trong\tdữ liệu.")
    doc.add_paragraph("1. Chuẩn bị dữ liệu đầu vào")
    doc.add_paragraph("2. Kiểm tra kết quả đầu ra")
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph("CHƯƠNG 2. KẾT QUẢ")
    doc.add_paragraph(
        "Phần thân bài tiếp theo, không thay chữ để làm tài liệu đẹp hơn."
    )
    table = doc.add_table(rows=3, cols=3)
    for cell, value in zip(table.rows[0].cells, ["Hạng mục", "Mô tả", "Giá trị"]):
        cell.text = value
        cell.paragraphs[0].runs[0].bold = True
    for cell, value in zip(
        table.rows[1].cells, ["A", "Mô tả dài hơn giúp tính độ rộng", "0012"]
    ):
        cell.text = value
    table.cell(2, 0).text = "B"
    table.cell(2, 2).text = "3,14"
    if case % 7 == 0:
        table.cell(2, 0).merge(table.cell(2, 1))
    if case % 4 == 0:
        doc.add_picture(BytesIO(PNG), width=Inches(7))
        doc.add_paragraph("Hình 2.1. Minh họa không được đổi nội dung")
    if case % 6 == 0:
        paragraph = doc.add_paragraph()
        math = OxmlElement("m:oMath")
        run = OxmlElement("m:r")
        text = OxmlElement("m:t")
        text.text = "E=mc²"
        run.append(text)
        math.append(run)
        paragraph._p.append(math)
    if case % 8 == 0:
        p = doc.add_paragraph("  if x:\n    return 42")
        p.runs[0].font.name = "Consolas"
    if case % 9 == 0:
        p = doc.add_paragraph()
        mark = OxmlElement("w:bookmarkStart")
        mark.set(q("id"), "77")
        mark.set(q("name"), "ProtectedBlank")
        p._p.append(mark)
    if case % 10 == 0:
        p = doc.add_paragraph("Không xóa revision: ")
        ins = OxmlElement("w:ins")
        ins.set(q("id"), "9")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "nội dung sửa"
        r.append(t)
        ins.append(r)
        p._p.append(ins)
    doc.add_heading("TÀI LIỆU THAM KHẢO", level=1)
    doc.add_paragraph("[1] Tài liệu tham khảo, không thay chuẩn trích dẫn.")
    doc.add_heading("PHỤ LỤC A", level=1)
    doc.add_paragraph("  Khoảng trắng của phụ lục cần được giữ.  ")
    if case % 11 == 0:
        section = doc.add_section()
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = (
            section.page_height,
            section.page_width,
        )
        doc.add_paragraph("Trang ngang giữ nguyên bố cục.")
    doc.sections[0].header.paragraphs[0].text = "Đầu trang gốc"
    doc.sections[0].footer.paragraphs[0].text = "Chân trang gốc"
    return save(doc)
