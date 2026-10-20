"""Contextual heading and cover alignment regressions using synthetic reports."""

from io import BytesIO
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches

from core.xmlutils import NS, q
from v1.analysis import Analysis, text_of
from v1.editor import Editor
from v1.package import Package, canonical
from v1_fixtures import PNG, save


def bold(doc, text):
    p = doc.add_paragraph(text)
    p.runs[0].bold = True
    return p


def analyze(doc):
    a = Analysis(Package(save(doc)))
    return a, a.public()["review_plan"]


def cover_doc():
    d = Document()
    bold(d, "TRƯỜNG ĐẠI HỌC MINH HỌA")
    bold(d, "BÁO CÁO MÔN HỌC")
    return d


class HierarchyTests(unittest.TestCase):
    def test_exercises_and_restarted_numeric_children_have_distinct_levels(self):
        d = Document()
        for exercise in range(1, 4):
            # Long exercise titles must not silently become ordinary body text.
            title = "BÀI TẬP " + str(exercise) + ": " + "THIẾT KẾ VÀ PHÂN TÍCH " * 9
            bold(d, title)
            for number in range(1, 4):
                bold(d, f"{number}. Nội dung cần trình bày")
                d.add_paragraph("Phần giải thích riêng của mục này.")
        a, plan = analyze(d)
        headings = [b for b in a.blocks if b["role"] == "heading"]
        self.assertEqual([b["level"] for b in headings], [1, 2, 2, 2] * 3)
        self.assertEqual(len(plan["automatic_headings"]), 12)
        self.assertFalse(plan["heading_groups"])
        editor = Editor(save(d), {"heading_levels": plan["automatic_headings"]}).run()
        self.assertTrue(editor.ledger.verify_bytes(editor.package.serialize()))
        roots = [b for b in editor.heading_blocks if b["level"] == 1]
        for block in roots[1:]:
            self.assertEqual(
                editor.analysis.elements[block["id"]]
                .find("w:pPr/w:pageBreakBefore", NS)
                .get(q("val")),
                "1",
            )

    def test_lettered_subsections_c_and_d_are_not_roman_major_sections(self):
        d = Document()
        d.add_heading("BÀI TẬP 1: PHÂN TÍCH", 1)
        bold(d, "1. Các giải pháp")
        for letter in "ABCD":
            bold(d, f"{letter}. Một giải pháp")
            d.add_paragraph("Nội dung giải thích giải pháp.")
        bold(d, "2. Tổng kết")
        d.add_paragraph("Đoạn tổng kết riêng.")
        a, plan = analyze(d)
        self.assertEqual(
            [b["level"] for b in a.blocks if b["role"] == "heading"],
            [1, 2, 3, 3, 3, 3, 2],
        )
        self.assertEqual(len(plan["automatic_headings"]), 6)

    def test_lowercase_letter_headings_have_the_same_parent(self):
        d = Document()
        d.add_heading("NỘI DUNG", 1)
        for number in (1, 2):
            bold(d, f"{number}. Một mục")
            for letter in "abc":
                bold(d, f"{letter}) Tiểu mục")
                d.add_paragraph("Giải thích tiểu mục riêng.")
        a, plan = analyze(d)
        self.assertEqual(
            [b["level"] for b in a.blocks if b["role"] == "heading"],
            [1] + [2, 3, 3, 3] * 2,
        )
        self.assertEqual(len(plan["automatic_headings"]), 8)

    def test_roman_sections_can_contain_restarted_single_numbers(self):
        d = Document()
        for roman in ("I", "II", "III"):
            bold(d, f"{roman}. PHẦN LỚN")
            for number in (1, 2):
                bold(d, f"{number}. Nội dung")
                d.add_paragraph("Giải thích riêng.")
        a, plan = analyze(d)
        self.assertEqual(
            [b["level"] for b in a.blocks if b["role"] == "heading"], [1, 2, 2] * 3
        )
        self.assertEqual(len(plan["automatic_headings"]), 9)

    def test_decimal_depth_depends_on_local_numbering_scheme(self):
        for single, expected in ((False, 2), (True, 3)):
            with self.subTest(single=single):
                d = Document()
                d.add_heading("CHƯƠNG 1. PHÂN TÍCH", 1)
                if single:
                    bold(d, "1. Tổng quan")
                for number in (1, 2):
                    bold(d, f"1.{number}. Một mục")
                    d.add_paragraph("Nội dung của mục nhiều cấp.")
                a, _ = analyze(d)
                self.assertEqual(
                    [
                        b["level"]
                        for b in a.blocks
                        if b["text"].startswith("1.") and ". " in b["text"][2:]
                    ],
                    [expected, expected],
                )

    def test_letter_children_of_roman_one_use_that_roman_parent(self):
        d = Document()
        for roman in ("I", "II"):
            bold(d, f"{roman}. PHẦN LỚN")
            for letter in "ab":
                bold(d, f"{letter}) Tiểu mục")
                d.add_paragraph("Giải thích riêng cho tiểu mục.")
        a, plan = analyze(d)
        self.assertEqual(
            [b["level"] for b in a.blocks if b["role"] == "heading"], [1, 2, 2] * 2
        )
        self.assertEqual(len(plan["automatic_headings"]), 6)

    def test_ambiguous_parent_does_not_automatically_promote_decimal_children(self):
        d = Document()
        for _ in range(2):
            bold(d, "BÀI TẬP 1: TRÙNG NHÃN")
            for number in (1, 2):
                bold(d, f"1.{number}. Nội dung")
                d.add_paragraph("Giải thích riêng.")
        _, plan = analyze(d)
        self.assertFalse(plan["automatic_headings"])

    def test_compact_letter_and_numeric_lists_stay_lists(self):
        d = Document()
        d.add_heading("NỘI DUNG", 1)
        for label in ("1.", "2.", "a)", "b)"):
            bold(d, f"{label} Việc cần làm")
        a, plan = analyze(d)
        self.assertFalse(plan["automatic_headings"])
        self.assertEqual([b["role"] for b in a.blocks[-2:]], ["manual_list"] * 2)

    def test_existing_word_levels_are_authoritative(self):
        d = Document()
        d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
        d.add_heading("1. Mục có chủ ý dùng cấp ba", 3)
        d.add_paragraph("Nội dung riêng.")
        bold(d, "2. Mục tiếp theo")
        d.add_paragraph("Nội dung khác.")
        a, _ = analyze(d)
        self.assertEqual(a.blocks[1]["level"], 3)
        self.assertEqual(a.blocks[1]["certainty"], "certain")

    def test_duplicate_exercise_labels_need_review(self):
        d = Document()
        for _ in range(2):
            bold(d, "BÀI TẬP 1: NỘI DUNG")
            d.add_paragraph("Phần trình bày riêng.")
        _, plan = analyze(d)
        self.assertFalse(plan["automatic_headings"])
        self.assertEqual(len(plan["heading_groups"]), 1)

    def test_ordinary_exercise_mentions_are_not_headings(self):
        d = Document()
        d.add_paragraph("Bài tập 1 được trình bày trong buổi học trước.")
        d.add_paragraph("Bài tập 2 cần được đọc cùng hướng dẫn sử dụng.")
        a, plan = analyze(d)
        self.assertTrue(all(b["role"] == "body" for b in a.blocks))
        self.assertFalse(plan["automatic_headings"])


class CoverAlignmentTests(unittest.TestCase):
    def test_interior_tabs_and_field_results_are_never_cleaned(self):
        d = cover_doc()
        p = d.add_paragraph("  Mã lớp:")
        p.add_run().add_tab()
        p.add_run("123  ")
        field_p = d.add_paragraph("  Ký hiệu: ")
        field = OxmlElement("w:fldSimple")
        field.set(q("instr"), " REF ExistingBookmark ")
        field_p._p.append(field)
        d.add_page_break()
        d.add_heading("NỘI DUNG", 1)
        data = save(d)
        e = Editor(data, {"mode": "layout", "cover": "academic"}).run()
        self.assertEqual(text_of(e.analysis.elements["p00003"]), "Mã lớp:123")
        self.assertEqual(
            len(e.analysis.elements["p00003"].findall(".//" + q("tab"))), 1
        )
        self.assertEqual(
            canonical(e.analysis.elements["p00004"]),
            canonical(Analysis(Package(data)).elements["p00004"]),
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_multiline_edges_are_cleaned_but_interior_spacing_is_preserved(self):
        d = cover_doc()
        p = d.add_paragraph("   Mã lớp:  123   ")
        p.add_run().add_tab()
        p.add_run().add_break()
        p.add_run().add_tab()
        p.add_run("   Người hướng dẫn:  A\u00a0B   ")
        d.add_page_break()
        d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
        e = Editor(save(d), {"mode": "layout", "cover": "academic"}).run()
        p = e.analysis.elements["p00003"]
        self.assertEqual(text_of(p), "Mã lớp:  123Người hướng dẫn:  A\u00a0B")
        self.assertEqual(len(p.findall(".//" + q("br"))), 1)
        self.assertFalse(p.findall(".//" + q("tab")))
        self.assertEqual(p.find("w:pPr/w:jc", NS).get(q("val")), "center")
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_centering_removes_character_and_directional_indents(self):
        d = cover_doc()
        p = d.add_paragraph("Tên nhóm: Minh họa")
        ind = OxmlElement("w:ind")
        for name in (
            "left",
            "right",
            "firstLine",
            "startChars",
            "endChars",
            "hangingChars",
        ):
            ind.set(q(name), "500")
        p._p.get_or_add_pPr().append(ind)
        d.add_page_break()
        d.add_heading("NỘI DUNG", 1)
        e = Editor(save(d), {"cover": "academic"}).run()
        ind = e.analysis.elements["p00003"].find("w:pPr/w:ind", NS)
        self.assertEqual(
            dict(ind.attrib), {q("left"): "0", q("right"): "0", q("firstLine"): "0"}
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_cover_logo_does_not_keep_an_old_first_line_indent(self):
        d = cover_doc()
        d.add_picture(BytesIO(PNG))
        d.paragraphs[-1].paragraph_format.first_line_indent = Inches(1)
        d.add_page_break()
        d.add_heading("NỘI DUNG", 1)
        e = Editor(save(d), {"cover": "academic"}).run()
        p = e.analysis.elements["p00003"]
        self.assertEqual(p.find("w:pPr/w:ind", NS).get(q("firstLine")), "0")
        self.assertEqual(p.find("w:pPr/w:jc", NS).get(q("val")), "center")
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_cover_table_is_centered_and_uses_compact_cell_spacing(self):
        d = cover_doc()
        t = d.add_table(rows=3, cols=3)
        for row in t.rows:
            for cell in row.cells:
                cell.text = "Thông tin"
        ind = OxmlElement("w:tblInd")
        ind.set(q("w"), "720")
        ind.set(q("type"), "dxa")
        t._tbl.tblPr.append(ind)
        d.add_page_break()
        d.add_heading("NỘI DUNG", 1)
        e = Editor(save(d), {"cover": "academic"}).run()
        table = e.analysis.elements["t00001"]
        self.assertEqual(table.find("w:tblPr/w:tblInd", NS).get(q("w")), "0")
        self.assertEqual(table.find("w:tblPr/w:jc", NS).get(q("val")), "center")
        self.assertEqual(
            set(table.xpath(".//w:spacing/@w:before", namespaces=NS)), {"40"}
        )
        self.assertEqual(set(table.xpath(".//w:jc/@w:val", namespaces=NS)), {"center"})
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_keep_cover_and_explicit_locks_preserve_exact_xml(self):
        d = cover_doc()
        d.add_paragraph("    Tên nhóm: A    ")
        d.add_page_break()
        d.add_heading("NỘI DUNG", 1)
        data = save(d)
        for options in (
            {"cover": "keep"},
            {"cover": "academic", "mode": "layout", "locked_blocks": ["p00003"]},
        ):
            e = Editor(data, options).run()
            original = Analysis(Package(data))
            self.assertEqual(
                canonical(e.analysis.elements["p00003"]),
                canonical(original.elements["p00003"]),
            )
            self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_strict_mode_preserves_padding_characters(self):
        d = cover_doc()
        d.add_paragraph("    Tên nhóm: A    ")
        d.add_page_break()
        d.add_heading("NỘI DUNG", 1)
        e = Editor(save(d), {"mode": "strict", "cover": "academic"}).run()
        self.assertEqual(text_of(e.analysis.elements["p00003"]), "    Tên nhóm: A    ")


if __name__ == "__main__":
    unittest.main()
