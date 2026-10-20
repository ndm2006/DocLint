from io import BytesIO
import unittest
from docx import Document
from docx.oxml import OxmlElement
from v1.analysis import Analysis
from v1.editor import Editor
from v1.package import Package
from v1_fixtures import PNG, save


def plan(doc):
    return Analysis(Package(save(doc))).public()["review_plan"]


class GuidedReviewTests(unittest.TestCase):
    def test_plain_text_established_headings_and_unique_caption_need_no_choices(self):
        d = Document()
        d.add_heading("Existing section", 1)
        for _ in range(20):
            d.add_paragraph(
                "This is an ordinary body paragraph that should not need review."
            )
        table = d.add_table(rows=2, cols=2)
        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        d.add_picture(BytesIO(PNG))
        d.add_paragraph("Hình 1. Minh họa")
        result = plan(d)
        self.assertEqual(result["question_count"], 0)
        self.assertEqual(result["counts"]["existing_headings"], 1)
        self.assertEqual(result["counts"]["linked_captions"], 1)

    def test_twenty_ambiguous_numbered_lines_form_one_question(self):
        d = Document()
        for index in range(1, 21):
            d.add_paragraph(f"{index}. Có thể là danh sách hoặc đề mục").runs[
                0
            ].bold = True
        result = plan(d)
        self.assertFalse(result["automatic_headings"])
        self.assertEqual(len(result["heading_groups"]), 1)
        self.assertEqual(len(result["heading_groups"][0]["block_ids"]), 20)
        self.assertEqual(result["heading_groups"][0]["default"], "keep")
        original = save(d)
        ids = result["heading_groups"][0]["block_ids"]
        editor = Editor(original, {"preserved_candidates": ids}).run()
        self.assertNotIn("heading_review", [w["code"] for w in editor.warnings])
        self.assertTrue(editor.ledger.verify_bytes(editor.package.serialize()))

    def test_consistent_chapters_with_separate_body_text_are_automatic(self):
        d = Document()
        for index in range(1, 4):
            d.add_paragraph(f"CHƯƠNG {index}. MỤC TIÊU").runs[0].bold = True
            d.add_paragraph("Đây là phần nội dung của chương được trình bày riêng.")
        result = plan(d)
        self.assertEqual(len(result["automatic_headings"]), 3)
        self.assertFalse(result["heading_groups"])

    def test_repeated_or_out_of_order_chapter_numbers_are_not_automatic(self):
        for numbers in ([3, 1, 2], [1, 1, 2]):
            with self.subTest(numbers=numbers):
                d = Document()
                for index, number in enumerate(numbers):
                    d.add_paragraph(f"CHƯƠNG {number}. MỤC {index}").runs[0].bold = True
                    d.add_paragraph("Phần nội dung sau tiêu đề, không thay chữ.")
                result = plan(d)
                self.assertFalse(result["automatic_headings"])
                self.assertEqual(len(result["heading_groups"]), 1)

    def test_only_ambiguous_manual_spacing_is_grouped(self):
        d = Document()
        for _ in range(8):
            d.add_paragraph(
                "    "
                + "Thân bài thông thường khá dài, có thể làm sạch thụt đầu dòng theo mẫu. "
                * 3
            )
        d.add_paragraph("              Dòng cần xác nhận căn lề")
        d.add_paragraph("              Dòng thứ hai cùng cách căn")
        result = plan(d)
        self.assertEqual(len(result["layout_groups"]), 1)
        self.assertEqual(len(result["layout_groups"][0]["block_ids"]), 2)

    def test_cover_boundary_stops_before_introductory_prose(self):
        d = Document()
        d.add_heading("BÁO CÁO MÔN HỌC", 1)
        for _ in range(8):
            d.add_paragraph()
        intro = d.add_paragraph(
            "Đây là đoạn mở đầu thân bài và không phải thông tin bìa."
        )
        d.add_heading("CHƯƠNG 1. TỔNG QUAN", 1)
        a = Analysis(Package(save(d)))
        self.assertEqual(a.cover_end, "p00001")
        self.assertNotEqual(a.elements[a.cover_end].xpath("string(.)"), intro.text)
        result = a.public()["review_plan"]
        self.assertEqual(result["cover_candidates"], ["p00001"])
        self.assertFalse(
            any("p00001" in g["block_ids"] for g in result["heading_groups"])
        )

    def test_confirming_headers_is_one_question_for_similar_tables(self):
        d = Document()
        for _ in range(3):
            table = d.add_table(rows=2, cols=2)
            for cell in table.rows[0].cells:
                cell.text = "Tên cột"
                cell.paragraphs[0].runs[0].bold = True
        result = plan(d)
        self.assertEqual(len(result["table_groups"]), 1)
        self.assertEqual(len(result["table_groups"][0]["block_ids"]), 3)

    def test_compact_report_title_is_cover_but_first_chapter_is_not(self):
        d = Document()
        d.add_heading("BÁO CÁO MÔN HỌC", 1)
        d.add_heading("TỔNG QUAN", 1)
        d.add_paragraph("Nội dung báo cáo cần giữ nguyên.")
        self.assertEqual(Analysis(Package(save(d))).cover_end, "p00001")
        d = Document()
        for index in range(1, 3):
            d.add_heading(f"CHƯƠNG {index}. TỔNG QUAN", 1)
            d.add_paragraph("Báo cáo này trình bày nội dung cần được giữ nguyên.")
        self.assertIsNone(Analysis(Package(save(d))).cover_end)
        d = Document()
        d.add_paragraph(
            "Đây là phần thân bài có nhắc đến báo cáo để giới thiệu nội dung."
        )
        d.add_paragraph("Các đoạn văn thường không được biến thành bìa.")
        d.add_paragraph()
        d.add_heading("Phần nội dung", 1)
        self.assertIsNone(Analysis(Package(save(d))).cover_end)
