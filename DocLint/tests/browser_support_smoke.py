"""A renderer failure must leave the review and verified Word download usable."""

from io import BytesIO
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
from urllib.error import URLError
from urllib.request import urlopen

from docx import Document
from playwright.sync_api import sync_playwright, expect

URL = "http://127.0.0.1:18083"
SERVER = """
import sys
from unittest.mock import patch
import uvicorn
from api import create_app
with patch('v1.renderer._safe_preview_copy', side_effect=PermissionError('test-only-private-marker')):
    uvicorn.run(create_app(sys.argv[1]), host='127.0.0.1', port=18083, log_level='error', access_log=False)
"""


def main():
    with tempfile.TemporaryDirectory() as tmp:
        process = subprocess.Popen(
            [sys.executable, "-c", SERVER, str(Path(tmp) / "store")],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
        try:
            for _ in range(100):
                try:
                    with urlopen(URL + "/api/v1/health", timeout=1):
                        break
                except URLError:
                    if process.poll() is not None:
                        raise AssertionError("Support test server did not start")
                    time.sleep(0.1)
            doc = Document()
            doc.add_paragraph("Nội dung thử nghiệm được giữ nguyên.")
            data = BytesIO()
            doc.save(data)
            with sync_playwright() as p:
                browser_path = os.environ.get("DOCLINT_BROWSER_PATH")
                browser = p.chromium.launch(
                    headless=True,
                    args=["--no-sandbox"],
                    **({"executable_path": browser_path} if browser_path else {}),
                )
                try:
                    page = browser.new_page()
                    page.set_default_timeout(30000)
                    errors = []
                    page.on("pageerror", lambda error: errors.append(str(error)))
                    page.goto(URL)
                    page.locator("#file-input").set_input_files(
                        {
                            "name": "private-test.docx",
                            "mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            "buffer": data.getvalue(),
                        }
                    )
                    page.locator("#submit-button").click()
                    expect(page.locator("#review")).to_be_visible()
                    expect(page.locator("#diagnostic-button")).to_be_visible()
                    expect(page.locator("#original-review-note")).to_contain_text(
                        "Mã chẩn đoán"
                    )
                    with page.expect_download() as pending:
                        page.locator("#diagnostic-button").click()
                    packet = json.loads(
                        Path(pending.value.path()).read_text(encoding="utf-8")
                    )
                    assert packet["phase"] == "render_pdf"
                    assert packet["exception"] == "PermissionError"
                    assert "private-test" not in json.dumps(packet)
                    page.locator("#format-button").click()
                    expect(page.locator("#results")).to_be_visible()
                    expect(page.locator("#diagnostic-button")).to_be_visible()
                    with page.expect_download() as pending:
                        page.locator("#download-button").click()
                    result = Document(pending.value.path())
                    assert result.paragraphs[0].text == doc.paragraphs[0].text
                    page.reload()
                    expect(page.locator("#results")).to_be_visible()
                    expect(page.locator("#diagnostic-button")).to_be_visible()
                    page.locator("#delete-session").click()
                    expect(page.locator("#upload-workspace")).to_be_visible()
                    expect(page.locator("#diagnostic-tools")).to_be_hidden()
                    assert not errors, errors
                    print(
                        "PASS: preview failure -> review -> safe diagnostic download -> verified Word -> reload -> delete; no JS errors."
                    )
                finally:
                    browser.close()
        finally:
            process.terminate()
            try:
                process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait()


if __name__ == "__main__":
    main()
