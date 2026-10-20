"""Regression coverage for the prioritized report-layout review."""

from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from docx import Document
from docx.oxml import OxmlElement
from docx.shared import Inches
from PIL import Image
from core.processor import DocumentError
from core.xmlutils import NS, q
from v1.analysis import Analysis, text_of
from v1.editor import Editor
from v1.package import Package, canonical
from v1.pipeline import format_document
from v1.renderer import find_office
from v1_fixtures import PNG, fixture, save
from test_v1_frontmatter import sample, page_texts


def inline_prose():
    d = Document()
    p = d.add_paragraph("FR01. Giữ nguyên yêu cầu có tên hàm ", style="List Number")
    p.add_run("ZipFile").font.name = "Consolas"
    p.add_run(" trong đoạn mô tả thông thường.")
    p.paragraph_format.line_spacing = 1.25
    return save(d)


def objects():
    d = Document()
    d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
    d.add_paragraph("Nội dung tham chiếu Hình 1.2 và Bảng 1.1 giữ nguyên.")
    for number in (1, 2):
        d.add_picture(BytesIO(PNG), width=Inches(1))
        d.add_paragraph(f"Hình 1.{number}. Chú thích hình có sẵn")
        d.add_paragraph(
            f"Nội dung sau hình thứ {number} để xác định liên kết duy nhất."
        )
    d.add_paragraph("Bảng 1.1. Dữ liệu có sẵn")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Tên"
    table.cell(0, 1).text = "Giá trị"
    table.cell(1, 0).text = "A"
    table.cell(1, 1).text = "0012"
    return save(d)


class PriorityLayoutTests(unittest.TestCase):
    def test_inline_identifier_does_not_protect_a_whole_numbered_requirement(self):
        data = inline_prose()
        a = Analysis(Package(data))
        self.assertEqual(a.blocks[0]["role"], "list")
        editor = Editor(data).run()
        p = editor.package.document.find(".//" + q("p"))
        self.assertEqual(p.find("w:pPr/w:spacing", NS).get(q("line")), "360")
        self.assertEqual(
            set(p.xpath(".//w:rFonts/@w:ascii", namespaces=NS)), {"Times New Roman"}
        )
        self.assertEqual(text_of(p), text_of(Package(data).document))
        self.assertIsNotNone(p.find("w:pPr/w:pStyle", NS))
        self.assertTrue(editor.ledger.verify_bytes(editor.package.serialize()))

    def test_prose_entirely_in_consolas_is_still_prose(self):
        d = Document()
        d.add_paragraph("Đây là đoạn văn được sao chép với phông đơn cách.").runs[
            0
        ].font.name = "Consolas"
        self.assertEqual(Analysis(Package(save(d))).blocks[0]["role"], "body")

    def test_backmatter_normalizes_without_altering_characters(self):
        d = Document()
        d.add_heading("TÀI LIỆU THAM KHẢO", 1)
        d.add_paragraph(" [1] Giữ nguyên cách ghi nguồn, 0012. ")
        d.add_heading("PHỤ LỤC A", 1)
        d.add_paragraph("  Phụ lục giữ cả khoảng trắng đầu.  ")
        data = save(d)
        e = Editor(data, {"mode": "layout"}).run()
        self.assertEqual(
            e.package.document.xpath(".//w:t/text()", namespaces=NS),
            Package(data).document.xpath(".//w:t/text()", namespaces=NS),
        )
        for bid in ("p00001", "p00003"):
            p = e.analysis.elements[bid]
            self.assertEqual(p.find("w:pPr/w:outlineLvl", NS).get(q("val")), "0")
            self.assertEqual(p.find("w:pPr/w:pageBreakBefore", NS).get(q("val")), "1")
            self.assertEqual(p.find(".//w:color", NS).get(q("val")), "000000")
        self.assertEqual(
            e.analysis.elements["p00002"].find("w:pPr/w:spacing", NS).get(q("line")),
            "360",
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))
        locked = Editor(
            data, {"preserve_references": True, "preserve_appendices": True}
        ).run()
        for block in locked.analysis.blocks:
            self.assertEqual(
                canonical(locked.analysis.elements[block["id"]]),
                canonical(Analysis(Package(data)).elements[block["id"]]),
            )

    def test_shared_page_geometry_is_not_blocked_by_automatic_cover_lock(self):
        data = sample()
        e = Editor(data).run()
        self.assertTrue(e.locked)
        self.assertEqual(
            e.package.document.find(".//" + q("pgSz")).get(q("w")), "11906"
        )
        e = Editor(data, {"locked_blocks": ["p00001"]}).run()
        self.assertEqual(
            canonical(e.package.document.find(".//" + q("sectPr"))),
            canonical(Package(data).document.find(".//" + q("sectPr"))),
        )

    def test_lead_in_is_not_a_heading_question_and_stays_with_list(self):
        d = Document()
        d.add_paragraph("Hạn chế kỹ thuật của giải pháp cũ:").runs[0].bold = True
        d.add_paragraph("Điểm hạn chế cần giữ nội dung.", style="List Bullet")
        data = save(d)
        a = Analysis(Package(data))
        self.assertEqual(a.blocks[0]["role"], "lead_in")
        self.assertEqual(a.public()["review_plan"]["question_count"], 0)
        e = Editor(data).run()
        self.assertEqual(
            e.analysis.elements["p00001"].find("w:pPr/w:keepNext", NS).get(q("val")),
            "1",
        )

    def test_short_table_keeps_its_rows_together_but_releases_last_row(self):
        d = Document()
        t = d.add_table(rows=4, cols=2)
        for i, row in enumerate(t.rows):
            row.cells[0].text = f"Hàng {i}"
            row.cells[1].text = "Nội dung"
        e = Editor(save(d)).run()
        rows = e.package.document.findall(".//" + q("tr"))
        for row in rows[:-1]:
            self.assertEqual(
                set(row.xpath(".//w:keepNext/@w:val", namespaces=NS)), {"1"}
            )
        self.assertEqual(
            set(rows[-1].xpath(".//w:keepNext/@w:val", namespaces=NS)), {"0"}
        )
        second = Editor(e.package.serialize()).run()
        self.assertEqual(second.changes, [])

    def test_cover_note_is_not_a_report_title(self):
        d = Document()
        d.add_paragraph("TRƯỜNG ĐẠI HỌC MINH HỌA")
        d.add_paragraph("BÁO CÁO MÔN HỌC")
        note = "Đây là tài liệu tổng hợp để thử phần mềm và không phải báo cáo của một người dùng thật."
        d.add_paragraph(note).runs[0].bold = True
        d.add_page_break()
        d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
        e = Editor(save(d), {"mode": "layout", "cover": "academic"}).run()
        p = e.analysis.elements["p00003"]
        self.assertEqual(p.find(".//w:sz", NS).get(q("val")), "22")
        self.assertEqual(p.find(".//w:b", NS).get(q("val")), "0")
        self.assertEqual(text_of(p), note)

    def test_code_styling_preserves_indentation_and_control_characters(self):
        d = Document()
        d.add_paragraph("  if x:\n    return 42").runs[0].font.name = "Consolas"
        data = save(d)
        e = Editor(data, {"mode": "layout"}).run()
        self.assertEqual(e.analysis.blocks[0]["role"], "code")
        p = e.analysis.elements["p00001"]
        self.assertEqual(p.find("w:pPr/w:shd", NS).get(q("fill")), "F3F5F4")
        self.assertEqual(
            e.package.document.xpath(".//w:t/text()", namespaces=NS),
            Package(data).document.xpath(".//w:t/text()", namespaces=NS),
        )
        self.assertEqual(len(p.findall(".//" + q("br"))), 1)
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_high_resolution_diagram_enlarges_but_tiny_image_does_not(self):
        raw = BytesIO()
        Image.new("RGB", (1920, 1080), "white").save(raw, format="PNG")
        d = Document()
        d.add_picture(BytesIO(raw.getvalue()), width=Inches(1))
        d.add_picture(BytesIO(PNG), width=Inches(1))
        data = save(d)
        e = Editor(data).run()
        extents = e.package.document.xpath(
            ".//wp:extent",
            namespaces={
                "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
            },
        )
        self.assertGreater(int(extents[0].get("cx")), 914400 * 3)
        self.assertEqual(int(extents[1].get("cx")), 914400)
        self.assertAlmostEqual(
            int(extents[0].get("cx")) / int(extents[0].get("cy")), 1920 / 1080, places=4
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_seq_and_ref_fields_preserve_every_caption_and_reference_character(self):
        data = objects()
        e = Editor(data, {"caption_fields": True}).run()
        instructions = e.package.document.xpath(".//w:instrText/text()", namespaces=NS)
        self.assertEqual(sum("SEQ " in v for v in instructions), 3)
        self.assertEqual(sum("REF " in v for v in instructions), 2)
        self.assertEqual(text_of(e.package.document), text_of(Package(data).document))
        names = set(
            e.package.document.xpath(".//w:bookmarkStart/@w:name", namespaces=NS)
        )
        for value in instructions:
            if value.strip().startswith("REF "):
                self.assertIn(value.split()[1], names)
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))
        # Cached field text is also part of the integrity gate.
        e.package.document.xpath(
            './/w:instrText[contains(text(),"SEQ")]/../following-sibling::w:r[w:t][1]/w:t',
            namespaces=NS,
        )[0].text = "999"
        with self.assertRaises(DocumentError):
            e.ledger.verify_bytes(e.package.serialize())

    def test_nonconsecutive_caption_numbers_are_not_renumbered(self):
        data = objects()
        p = Package(data)
        node = next(
            n for n in p.document.iter(q("t")) if (n.text or "").startswith("Hình 1.2")
        )
        node.text = node.text.replace("1.2", "1.4")
        data = p.serialize()
        e = Editor(data, {"caption_fields": True}).run()
        instructions = e.package.document.xpath(".//w:instrText/text()", namespaces=NS)
        self.assertFalse(any("DLFigure" in v for v in instructions))
        self.assertEqual(text_of(e.package.document), text_of(Package(data).document))
        self.assertIn("caption_sequence_ambiguous", [w["code"] for w in e.warnings])

    def test_native_list_preserves_start_number_and_style_punctuation(self):
        d = Document()
        d.add_paragraph("3) Mục thứ ba")
        d.add_paragraph("4) Mục thứ tư")
        data = save(d)
        e = Editor(data, {"mode": "layout", "convert_manual_lists": True}).run()
        self.assertEqual(
            len(
                e.package.document.xpath(
                    ".//w:pPr/w:numPr/w:numId/@w:val", namespaces=NS
                )
            ),
            2,
        )
        root = e.package.xml("word/numbering.xml")
        self.assertEqual(
            root.xpath("./w:abstractNum[last()]/w:lvl/w:start/@w:val", namespaces=NS),
            ["3"],
        )
        self.assertEqual(
            root.xpath("./w:abstractNum[last()]/w:lvl/w:lvlText/@w:val", namespaces=NS),
            ["%1)"],
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))
        strict = Editor(data, {"convert_manual_lists": True}).run()
        self.assertEqual(
            text_of(strict.package.document), text_of(Package(data).document)
        )
        self.assertFalse(strict.package.document.xpath(".//w:numPr", namespaces=NS))

    def test_existing_manual_toc_separates_pages_without_rewriting_entries(self):
        d = Document()
        d.add_heading("BÁO CÁO MÔN HỌC", 1)
        d.add_paragraph("MỤC LỤC")
        d.add_paragraph("CHƯƠNG 1. NỘI DUNG\t7")
        d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
        d.add_paragraph("Đoạn nội dung gốc.")
        data = save(d)
        e = Editor(data).run()
        self.assertEqual(text_of(e.package.document), text_of(Package(data).document))
        self.assertEqual(
            e.analysis.elements["p00002"]
            .find("w:pPr/w:pageBreakBefore", NS)
            .get(q("val")),
            "1",
        )
        self.assertEqual(
            e.analysis.elements["p00004"]
            .find("w:pPr/w:pageBreakBefore", NS)
            .get(q("val")),
            "1",
        )

    def test_special_sections_and_explicit_locks_do_not_get_academic_restarts(self):
        e = Editor(fixture(0), {"page_numbers": "academic"}).run()
        self.assertFalse(e.numbering_applied)
        self.assertIn("academic_sections_preserved", [w["code"] for w in e.warnings])
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    @unittest.skipUnless(find_office(), "LibreOffice is required")
    def test_academic_cover_roman_front_body_one_and_appendix_continues(self):
        source = Package(sample(True))
        body = source.document.find(q("body"))
        appendix = OxmlElement("w:p")
        pp = OxmlElement("w:pPr")
        style = OxmlElement("w:pStyle")
        style.set(q("val"), "Heading1")
        pp.append(style)
        appendix.append(pp)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "PHỤ LỤC A"
        run.append(text)
        appendix.append(run)
        body.insert(len(body) - 1, appendix)
        with tempfile.TemporaryDirectory() as tmp:
            out, r = format_document(
                source.serialize(),
                {
                    "cover": "academic",
                    "toc": "insert",
                    "mode": "layout",
                    "page_numbers": "academic",
                },
                tmp,
            )
            self.assertTrue(r["integrity"]["verified"])
            self.assertEqual(r["toc_status"], "verified")
            self.assertEqual(r["status"], "ready")
            p = Package(out)
            self.assertEqual(
                p.document.xpath(".//w:sectPr/w:pgNumType/@w:fmt", namespaces=NS),
                ["decimal", "lowerRoman", "decimal", "decimal"],
            )
            self.assertEqual(
                p.document.xpath(".//w:sectPr/w:pgNumType/@w:start", namespaces=NS),
                ["1", "1", "1"],
            )
            pages = page_texts(Path(tmp) / "result.pdf")
            self.assertEqual(len(pages), 4)
            self.assertTrue(pages[1].startswith("i "))
            self.assertTrue(pages[2].startswith("1 "))
            self.assertTrue(pages[3].startswith("2 "))
            self.assertEqual(
                p.document.xpath(".//w:sectPr/w:pgBorders/@w:display", namespaces=NS),
                ["firstPage"],
            )

    @unittest.skipUnless(find_office(), "LibreOffice is required")
    def test_object_lists_resolve_repeated_caption_text_to_body_pages(self):
        data = objects()
        with tempfile.TemporaryDirectory() as tmp:
            output, report = format_document(
                data,
                {
                    "toc": "insert",
                    "caption_fields": True,
                    "object_lists": True,
                    "page_numbers": "academic",
                },
                tmp,
            )
            self.assertEqual(report["status"], "ready", report["warnings"])
            self.assertEqual(report["toc_status"], "verified")
            p = Package(output)
            lists = p.document.xpath(
                './/w:sdt[w:sdtPr/w:tag[@w:val="DocLintObjectList"]]', namespaces=NS
            )
            self.assertEqual(len(lists), 2)
            for container in lists:
                self.assertNotIn("…", container.xpath(".//w:t/text()", namespaces=NS))
                self.assertEqual(
                    container.xpath(
                        './/w:fldChar[@w:fldCharType="begin"]/@w:dirty', namespaces=NS
                    ),
                    ["false"],
                )
            pages = page_texts(Path(tmp) / "result.pdf")
            self.assertIn("DANH MỤC HÌNH", pages[1])
            self.assertIn("DANH MỤC BẢNG", pages[2])
            self.assertTrue(pages[3].startswith("1 "))
            self.assertEqual(
                p.document.xpath(".//w:updateFields/@w:val", namespaces=NS), []
            )
            self.assertEqual(
                p.xml("word/settings.xml").xpath(
                    "./w:updateFields/@w:val", namespaces=NS
                ),
                ["false"],
            )

    def test_glossary_uses_explicit_definitions_and_does_not_invent_meanings(self):
        d = Document()
        d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
        d.add_paragraph("API: Giao diện lập trình ứng dụng")
        d.add_paragraph("XML: Ngôn ngữ đánh dấu mở rộng")
        d.add_paragraph("DOCX và PDF chưa có định nghĩa trong nội dung này.")
        data = save(d)
        e = Editor(data, {"object_lists": True}).run()
        glossary = e.package.document.xpath(
            './/w:sdt[w:sdtPr/w:tag[@w:val="DocLintGlossary"]]', namespaces=NS
        )
        self.assertEqual(len(glossary), 1)
        rows = glossary[0].xpath(".//w:tr", namespaces=NS)
        self.assertEqual(len(rows), 3)
        self.assertIn("API", text_of(rows[1]))
        self.assertIn("XML", text_of(rows[2]))
        self.assertNotIn("DOCX", text_of(glossary[0]))
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_signature_relocation_requires_option_and_preserves_name_group(self):
        d = Document()
        d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
        d.add_paragraph("Sinh viên thực hiện")
        d.add_paragraph("Nguyễn Văn Minh")
        d.add_paragraph("Đây là phần thân bài còn lại cần đặt trước chữ ký.")
        d.add_heading("TÀI LIỆU THAM KHẢO", 1)
        data = save(d)
        kept = Editor(data).run()
        self.assertEqual(
            text_of(kept.package.document), text_of(Package(data).document)
        )
        moved = Editor(data, {"signature_position": "end_of_body"}).run()
        paragraphs = moved.package.document.findall(".//" + q("body") + "/" + q("p"))
        texts = [text_of(p) for p in paragraphs]
        self.assertLess(
            texts.index("Đây là phần thân bài còn lại cần đặt trước chữ ký."),
            texts.index("Sinh viên thực hiện"),
        )
        self.assertEqual(
            texts.index("Nguyễn Văn Minh"), texts.index("Sinh viên thực hiện") + 1
        )
        self.assertLess(
            texts.index("Nguyễn Văn Minh"), texts.index("TÀI LIỆU THAM KHẢO")
        )
        self.assertTrue(moved.ledger.verify_bytes(moved.package.serialize()))

    def test_character_style_symbol_font_is_not_replaced(self):
        from docx.enum.style import WD_STYLE_TYPE

        d = Document()
        style = d.styles.add_style("SymbolGlyph", WD_STYLE_TYPE.CHARACTER)
        style.font.name = "Symbol"
        p = d.add_paragraph("Ký hiệu: ")
        p.add_run("a", style="SymbolGlyph")
        e = Editor(save(d)).run()
        run = e.package.document.xpath(
            './/w:r[w:rPr/w:rStyle[@w:val="SymbolGlyph"]]', namespaces=NS
        )[0]
        self.assertIsNone(run.find("w:rPr/w:rFonts", NS))
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    @unittest.skipUnless(find_office(), "LibreOffice is required")
    def test_kept_cover_manual_page_break_does_not_duplicate_section_break(self):
        with tempfile.TemporaryDirectory() as tmp:
            out, r = format_document(
                sample(True), {"toc": "insert", "page_numbers": "academic"}, tmp
            )
            self.assertTrue(r["integrity"]["verified"])
            self.assertEqual(r["status"], "ready", r["warnings"])
            self.assertEqual(len(page_texts(Path(tmp) / "result.pdf")), 3)
            self.assertEqual(
                Package(out).document.xpath(".//w:sdt//w:t/text()", namespaces=NS)[-1],
                "1",
            )

    @unittest.skipUnless(find_office(), "LibreOffice is required")
    def test_named_summary_is_roman_and_chapter_one_restarts_at_one(self):
        d = Document()
        d.add_heading("BÁO CÁO MÔN HỌC", 1)
        d.add_page_break()
        d.add_heading("TÓM TẮT NỘI DUNG BÁO CÁO", 1)
        d.add_paragraph("Phần tóm tắt có tiêu đề rõ ràng, giữ nguyên nội dung.")
        d.add_heading("CHƯƠNG 1. TỔNG QUAN", 1)
        d.add_paragraph("Nội dung chương một bắt đầu đánh số từ trang một.")
        with tempfile.TemporaryDirectory() as tmp:
            output, report = format_document(
                save(d),
                {
                    "cover": "academic",
                    "toc": "insert",
                    "page_numbers": "academic",
                    "mode": "layout",
                },
                tmp,
            )
            self.assertEqual(report["status"], "ready", report["warnings"])
            pages = page_texts(Path(tmp) / "result.pdf")
            self.assertEqual(len(pages), 4)
            self.assertTrue(pages[2].startswith("ii "))
            self.assertTrue(pages[3].startswith("1 "))
            toc = Package(output).document.xpath(".//w:sdt//w:t/text()", namespaces=NS)
            self.assertIn("ii", toc)
            self.assertIn("1", toc)

    def test_caption_group_is_not_partly_converted_after_a_split_label(self):
        p = Package(objects())
        caption = next(
            n for n in p.document.iter(q("p")) if text_of(n).startswith("Hình 1.1")
        )
        run = caption.find(q("r"))
        run.find(q("t")).text = "Hình "
        second = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "1.1. Chú thích hình có sẵn"
        second.append(text)
        caption.append(second)
        e = Editor(p.serialize(), {"caption_fields": True}).run()
        self.assertFalse(
            any(
                "DLFigure" in v
                for v in e.package.document.xpath(
                    ".//w:instrText/text()", namespaces=NS
                )
            )
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_nested_and_object_tables_are_preserved(self):
        d = Document()
        table = d.add_table(rows=1, cols=1)
        table.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Bảng lồng"
        data = save(d)
        e = Editor(data).run()
        self.assertEqual(
            canonical(e.package.document.find(".//" + q("tbl"))),
            canonical(Package(data).document.find(".//" + q("tbl"))),
        )
        self.assertIn("complex_table", [w["code"] for w in e.warnings])
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_confirmed_ordinary_candidate_still_gets_consistent_typography(self):
        d = Document()
        p = d.add_paragraph("1. Có thể là tiêu đề hoặc danh sách")
        p.runs[0].bold = True
        p.paragraph_format.line_spacing = 1.25
        e = Editor(save(d), {"preserved_candidates": ["p00001"]}).run()
        self.assertEqual(
            e.analysis.elements["p00001"].find("w:pPr/w:spacing", NS).get(q("line")),
            "360",
        )
        self.assertFalse(
            e.analysis.elements["p00001"].xpath("./w:pPr/w:outlineLvl", namespaces=NS)
        )
        self.assertNotIn("heading_review", [w["code"] for w in e.warnings])
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_glossary_is_not_duplicated_when_processing_an_export(self):
        d = Document()
        d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
        d.add_paragraph("API: Giao diện lập trình ứng dụng")
        d.add_paragraph("XML: Ngôn ngữ đánh dấu mở rộng")
        one = Editor(save(d), {"object_lists": True}).run()
        two = Editor(one.package.serialize(), {"object_lists": True}).run()
        self.assertEqual(
            len(
                two.package.document.xpath(
                    './/w:sdtPr/w:tag[@w:val="DocLintGlossary"]', namespaces=NS
                )
            ),
            1,
        )
        self.assertIsNone(two.analysis.toc_region)
        self.assertTrue(two.ledger.verify_bytes(two.package.serialize()))

    def test_even_odd_footers_are_rejected_before_inserting_sections(self):
        p = Package(sample())
        setting = p.xml("word/settings.xml")
        flag = OxmlElement("w:evenAndOddHeaders")
        flag.set(q("val"), "1")
        setting.insert(0, flag)
        p.write_roots.add("word/settings.xml")
        e = Editor(
            p.serialize(),
            {"cover": "academic", "toc": "insert", "page_numbers": "academic"},
        ).run()
        self.assertFalse(e.numbering_applied)
        self.assertEqual(len(e.package.document.xpath(".//w:sectPr", namespaces=NS)), 1)
        self.assertIn("numbering_preserved", [w["code"] for w in e.warnings])
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))
