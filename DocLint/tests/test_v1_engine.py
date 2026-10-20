from io import BytesIO
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile
from docx import Document
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches
from lxml import etree
from core.processor import DocumentError
from core.xmlutils import NS, q
from v1.analysis import Analysis
from v1.editor import Editor
from v1.package import Package, R, canonical, structural_signature
from v1.pipeline import format_document
from v1.renderer import PreviewError, find_office, page_png, _safe_preview_copy
from v1_fixtures import PNG, fixture, save


def approved(data):
    a = Analysis(Package(data))
    return {
        b["id"]: b["level"]
        for b in a.blocks
        if b["role"] == "heading" and b.get("level") and not b.get("region")
    }


class EngineV1Tests(unittest.TestCase):
    def test_thirty_synthetic_reports_strict_and_layout(self):
        for case in range(30):
            with self.subTest(case=case):
                data = fixture(case)
                options = {
                    "mode": "layout" if case % 2 else "strict",
                    "heading_levels": approved(data),
                    "preserve_cover": case % 3 == 0,
                }
                editor = Editor(data, options).run()
                result = editor.package.serialize()
                self.assertTrue(editor.ledger.verify_bytes(result))
                Document(BytesIO(result))
                self.assertEqual(
                    editor.package.document.xpath("//m:t/text()", namespaces=NS),
                    Package(data).document.xpath("//m:t/text()", namespaces=NS),
                )
                if options["mode"] == "strict":
                    self.assertEqual(
                        editor.package.document.xpath("//w:t/text()", namespaces=NS),
                        Package(data).document.xpath("//w:t/text()", namespaces=NS),
                    )
                with (
                    zipfile.ZipFile(BytesIO(data)) as before,
                    zipfile.ZipFile(BytesIO(result)) as after,
                ):
                    for name in before.namelist():
                        if name not in {"word/document.xml", "word/styles.xml"}:
                            self.assertEqual(before.read(name), after.read(name), name)

    def test_unapproved_heading_is_unchanged(self):
        doc = Document()
        p = doc.add_paragraph("1.1. Một tiêu đề có thể có")
        p.runs[0].bold = True
        data = save(doc)
        editor = Editor(data).run()
        self.assertEqual(
            canonical(Package(data).document.find(".//" + q("p"))),
            canonical(editor.package.document.find(".//" + q("p"))),
        )
        self.assertIn("heading_review", [w["code"] for w in editor.warnings])

    def test_approved_headings_have_outline_and_keep_next(self):
        data = fixture(1)
        editor = Editor(data, {"heading_levels": approved(data)}).run()
        for block in editor.heading_blocks:
            p = editor.analysis.elements[block["id"]]
            self.assertEqual(
                p.find("w:pPr/w:outlineLvl", NS).get(q("val")), str(block["level"] - 1)
            )
            self.assertEqual(p.find("w:pPr/w:keepNext", NS).get(q("val")), "1")
            self.assertEqual(
                "".join(p.xpath(".//w:t/text()", namespaces=NS)),
                "".join(editor.ledger.pairs[p].xpath(".//w:t/text()", namespaces=NS)),
            )

    def test_false_bold_is_not_a_heading(self):
        d = Document()
        p = d.add_paragraph("1. Chuẩn bị dữ liệu")
        r = p.runs[0]
        r.bold = False
        a = Analysis(Package(save(d)))
        self.assertEqual(a.blocks[0]["role"], "manual_list")

    def test_style_inherited_numbering_is_preserved(self):
        d = Document()
        d.add_paragraph("Mục danh sách", style="List Number")
        data = save(d)
        editor = Editor(data).run()
        p = editor.package.document.find(".//" + q("p"))
        self.assertEqual(p.find("w:pPr/w:pStyle", NS).get(q("val")), "ListNumber")
        self.assertEqual(
            structural_signature(Package(data).document),
            structural_signature(editor.package.document),
        )

    def test_locked_cover_reference_appendix_and_section(self):
        data = fixture(3)
        editor = Editor(data, {"heading_levels": approved(data)}).run()
        self.assertGreater(len(editor.locked), 0)
        for node, before, _ in editor.ledger.locks:
            self.assertEqual(canonical(node), before)
        original = Package(data)
        self.assertNotEqual(
            canonical(original.document.xpath(".//w:sectPr", namespaces=NS)[0]),
            canonical(editor.package.document.xpath(".//w:sectPr", namespaces=NS)[0]),
        )
        self.assertEqual(
            editor.package.document.find(".//" + q("pgSz")).get(q("w")), "11906"
        )

    def test_manual_block_lock(self):
        data = fixture(1)
        a = Analysis(Package(data))
        b = next(b for b in a.blocks if b["role"] == "body")
        editor = Editor(data, {"mode": "layout", "locked_blocks": [b["id"]]}).run()
        self.assertEqual(
            canonical(a.elements[b["id"]]), canonical(editor.analysis.elements[b["id"]])
        )
        self.assertTrue(editor.ledger.verify())

    def test_whitespace_consent_and_internal_spaces(self):
        d = Document()
        d.add_paragraph("    Đây  là  nội dung  gốc. ")
        data = save(d)
        strict = Editor(data).run()
        layout = Editor(data, {"mode": "layout"}).run()
        self.assertEqual(
            strict.package.document.xpath("//w:t/text()", namespaces=NS),
            ["    Đây  là  nội dung  gốc. "],
        )
        self.assertEqual(
            layout.package.document.xpath("//w:t/text()", namespaces=NS),
            ["Đây  là  nội dung  gốc. "],
        )
        self.assertTrue(layout.ledger.events)

    def test_tabs_in_body_and_code_are_never_removed(self):
        d = Document()
        d.add_paragraph("\tA\tB")
        p = d.add_paragraph("   if x:\n     return 1")
        p.runs[0].font.name = "Consolas"
        data = save(d)
        e = Editor(data, {"mode": "layout"}).run()
        self.assertEqual(
            Package(data).document.xpath("//w:t/text()", namespaces=NS),
            e.package.document.xpath("//w:t/text()", namespaces=NS),
        )
        self.assertEqual(len(e.package.document.xpath("//w:tab", namespaces=NS)), 2)

    def test_protected_blank_and_table_cell_are_not_removed(self):
        d = Document()
        d.add_paragraph()
        p = d.add_paragraph()
        m = OxmlElement("w:bookmarkStart")
        m.set(q("id"), "1")
        m.set(q("name"), "keep")
        p._p.append(m)
        d.add_paragraph()
        d.add_paragraph()
        d.add_paragraph("CHƯƠNG 2. NỘI DUNG")
        data = save(d)
        e = Editor(data, {"mode": "layout", "heading_levels": approved(data)}).run()
        self.assertEqual(
            len(e.package.document.xpath("//w:bookmarkStart", namespaces=NS)), 1
        )
        self.assertEqual(len(e.package.document.xpath("//w:p", namespaces=NS)), 3)
        self.assertEqual(
            e.package.document.find(".//w:pageBreakBefore", NS).get(q("val")), "1"
        )

    def test_idempotence_no_style_growth(self):
        data = fixture(1)
        e = Editor(data, {"heading_levels": approved(data)}).run()
        out = e.package.serialize()
        second = Editor(out).run()
        self.assertEqual(second.changes, [])
        self.assertEqual(len(second.package.styles), len(e.package.styles))

    def test_simple_table_width_header_and_cells(self):
        d = Document()
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Tên"
        t.cell(0, 1).text = "Số"
        t.cell(1, 0).text = "A"
        t.cell(1, 1).text = "0012"
        data = save(d)
        e = Editor(data, {"table_headers": ["t00001"]}).run()
        table = e.analysis.elements["t00001"]
        self.assertIsNotNone(table.find("w:tr/w:trPr/w:tblHeader", NS))
        self.assertEqual(
            e.package.document.xpath("//w:t/text()", namespaces=NS),
            Package(data).document.xpath("//w:t/text()", namespaces=NS),
        )
        widths = [int(c.get(q("w"))) for c in table.findall("w:tblGrid/w:gridCol", NS)]
        self.assertEqual(sum(widths), e._available_width(e.analysis.blocks[0]))

    def test_merged_table_formats_without_changing_cell_topology(self):
        d = Document()
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).merge(t.cell(0, 1))
        t.cell(0, 0).text = "Ô gộp"
        data = save(d)
        e = Editor(data).run()
        self.assertNotEqual(
            canonical(Package(data).document.find(".//" + q("tbl"))),
            canonical(e.package.document.find(".//" + q("tbl"))),
        )
        self.assertEqual(
            structural_signature(Package(data).document),
            structural_signature(e.package.document),
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_image_fit_retains_ratio_and_bytes(self):
        data = fixture(4)
        e = Editor(data).run()
        before = Package(data).document.xpath(
            ".//wp:extent",
            namespaces={
                **NS,
                "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
            },
        )[0]
        after = e.package.document.xpath(
            ".//wp:extent",
            namespaces={
                **NS,
                "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
            },
        )[0]
        self.assertLess(int(after.get("cx")), int(before.get("cx")))
        self.assertAlmostEqual(
            int(before.get("cx")) / int(before.get("cy")),
            int(after.get("cx")) / int(after.get("cy")),
            places=4,
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_caption_link_ambiguity_requires_choice(self):
        from v1_fixtures import PNG

        d = Document()
        d.add_picture(BytesIO(PNG))
        d.add_paragraph("Hình 1. Minh họa")
        d.add_picture(BytesIO(PNG))
        data = save(d)
        a = Analysis(Package(data))
        b = next(b for b in a.blocks if b["role"] == "caption")
        self.assertIsNone(b["target"])
        self.assertEqual(len(b["targets"]), 2)
        e = Editor(data, {"caption_links": {b["id"]: b["targets"][0]}}).run()
        self.assertTrue(e.ledger.verify())

    def test_invalid_options_fail_before_modifying(self):
        data = fixture(1)
        for options in (
            {"heading_levels": {"missing": 2}},
            {"heading_levels": {"p00001": 7}},
            {"locked_blocks": ["missing"]},
            {"cover": "academic"},
            {"toc": "replace"},
            {"caption_links": {"p00001": "t00001"}},
        ):
            with self.subTest(options=options), self.assertRaises(DocumentError):
                Editor(data, options)

    def test_integrity_detects_text_math_and_media_mutation(self):
        data = fixture(6)
        e = Editor(data).run()
        e.package.document.find(".//" + q("t")).text = "Bị đổi"
        with self.assertRaises(DocumentError):
            e.ledger.verify()
        e = Editor(data).run()
        e.package.document.xpath("//m:t", namespaces=NS)[0].text = "Sai công thức"
        with self.assertRaises(DocumentError):
            e.ledger.verify()
        data = fixture(4)
        e = Editor(data).run()
        name = next(n for n in e.package.parts if n.startswith("word/media/"))
        e.package.parts[name] = b"corrupt"
        with self.assertRaises(DocumentError):
            e.ledger.verify()

    def test_serialized_gate_detects_mutation(self):
        data = fixture(1)
        e = Editor(data).run()
        out = e.package.serialize()
        target = BytesIO()
        with zipfile.ZipFile(BytesIO(out)) as src, zipfile.ZipFile(target, "w") as dst:
            for name in src.namelist():
                dst.writestr(
                    name,
                    src.read(name).replace("Đây".encode(), "Sai".encode())
                    if name == "word/document.xml"
                    else src.read(name),
                )
        with self.assertRaises(DocumentError):
            e.ledger.verify_bytes(target.getvalue())

    def test_manual_toc_region_is_explicit(self):
        d = Document()
        d.add_paragraph("MỤC LỤC")
        d.add_paragraph("CHƯƠNG 1. GIỚI THIỆU")
        d.add_paragraph("1.1. Mục tiêu")
        d.add_paragraph("CHƯƠNG 1. GIỚI THIỆU")
        d.add_paragraph("1.1. Mục tiêu")
        d.add_paragraph("Nội dung thật.")
        data = save(d)
        a = Analysis(Package(data))
        self.assertEqual(
            a.toc_region, {"start": "p00001", "end": "p00003", "kind": "manual"}
        )
        self.assertEqual(a.blocks[1]["role"], "toc")
        e = Editor(data, {"toc": "replace", "heading_levels": approved(data)}).run()
        self.assertEqual(len(e.toc_entries), 2)
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_toc_generated_field_is_updateable_and_preserves_body(self):
        data = fixture(1)
        e = Editor(data, {"toc": "insert", "heading_levels": approved(data)}).run()
        self.assertTrue(e.toc_entries)
        self.assertTrue(
            e.package.document.xpath(
                './/w:instrText[contains(text(),"TOC")]', namespaces=NS
            )
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))
        with self.assertRaises(DocumentError):
            Editor(e.package.serialize(), {"toc": "insert"}).run()

    def test_footer_preserves_existing_text_and_headers(self):
        data = fixture(1)
        e = Editor(data, {"page_numbers": "arabic"}).run()
        out = e.package.serialize()
        e.ledger.verify_bytes(out)
        with zipfile.ZipFile(BytesIO(out)) as z:
            roots = [z.read(n) for n in z.namelist() if n.startswith("word/doclint_")]
            self.assertTrue(any("Chân trang gốc".encode() in n for n in roots))
            self.assertEqual(
                z.read("word/header1.xml"), e.ledger.original_parts["word/header1.xml"]
            )

    def test_cover_requires_confirmed_boundary_and_no_invention(self):
        data = fixture(3)
        a = Analysis(Package(data))
        self.assertIsNotNone(a.cover_end)
        e = Editor(data, {"cover": "academic", "cover_end": a.cover_end}).run()
        self.assertEqual(
            e.package.document.xpath("//w:t/text()", namespaces=NS),
            Package(data).document.xpath("//w:t/text()", namespaces=NS),
        )
        self.assertIn("cover", [c["rule"] for c in e.changes])

    def test_numbering_clones_footer_images_links_and_absolute_targets(self):
        doc = Document()
        doc.add_paragraph("Nội dung gốc")
        footer = doc.sections[0].footer
        p = footer.paragraphs[0]
        p.add_run("Logo và liên kết: ")
        p.add_run().add_picture(BytesIO(PNG), width=Inches(0.2))
        rid = footer.part.relate_to(
            "https://example.org/", RT.HYPERLINK, is_external=True
        )
        link = OxmlElement("w:hyperlink")
        link.set(f"{{{R}}}id", rid)
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "Trang chủ"
        run.append(text)
        link.append(run)
        p._p.append(link)
        package = Package(save(doc))
        rels = package.xml("word/_rels/document.xml.rels")
        for rel in rels:
            if rel.get("Type") == RT.FOOTER:
                rel.set("Target", "/word/footer1.xml")
        package.write_roots.add("word/_rels/document.xml.rels")
        data = package.serialize()
        e = Editor(data, {"page_numbers": "arabic"}).run()
        self.assertTrue(e.numbering_applied)
        out = e.package.serialize()
        e.ledger.verify_bytes(out)
        reopened = Document(BytesIO(out))
        for f in [reopened.sections[0].footer, reopened.sections[0].first_page_footer]:
            self.assertIn("Trang chủ", "".join(f._element.itertext()))
            self.assertEqual(len(f._element.xpath(".//w:drawing")), 1)
            self.assertTrue(
                any(
                    r.target_ref == "https://example.org/" for r in f.part.rels.values()
                )
            )
            self.assertTrue(
                any(
                    r.reltype == RT.IMAGE and r.target_part.blob == PNG
                    for r in f.part.rels.values()
                )
            )
        with zipfile.ZipFile(BytesIO(out)) as z:
            self.assertEqual(
                z.read("word/footer1.xml"), Package(data).parts["word/footer1.xml"]
            )

    def test_nested_footer_fields_leave_numbering_unchanged(self):
        doc = Document()
        doc.add_paragraph("Văn bản")
        p = doc.sections[0].footer.paragraphs[0]._p
        for kind, instruction in [
            ("begin", " IF "),
            ("begin", " PAGE "),
            ("end", None),
            ("end", None),
        ]:
            run = OxmlElement("w:r")
            marker = OxmlElement("w:fldChar")
            marker.set(q("fldCharType"), kind)
            run.append(marker)
            p.append(run)
            if instruction:
                run = OxmlElement("w:r")
                node = OxmlElement("w:instrText")
                node.text = instruction
                run.append(node)
                p.append(run)
        data = save(doc)
        e = Editor(data, {"page_numbers": "arabic"}).run()
        self.assertFalse(e.numbering_applied)
        self.assertIn("numbering_preserved", [w["code"] for w in e.warnings])
        out = e.package.serialize()
        e.ledger.verify_bytes(out)
        self.assertEqual(
            Package(out).parts["word/footer1.xml"],
            Package(data).parts["word/footer1.xml"],
        )
        self.assertFalse(
            any(n.startswith("word/doclint_footer") for n in Package(out).parts)
        )

    def test_missing_footer_target_does_not_replace_footer(self):
        package = Package(fixture(1))
        name = "word/_rels/document.xml.rels"
        rels = package.xml(name)
        for rel in rels:
            if rel.get("Type") == RT.FOOTER:
                rel.set("Target", "missing-footer.xml")
        package.write_roots.add(name)
        data = package.serialize()
        e = Editor(data, {"page_numbers": "arabic"}).run()
        self.assertFalse(e.numbering_applied)
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))
        self.assertEqual(
            e.package.xml(name).xpath("string(*[contains(@Type,'/footer')]/@Target)"),
            "missing-footer.xml",
        )

    def test_numbering_does_not_activate_unused_first_page_content(self):
        doc = Document()
        doc.add_paragraph("Văn bản")
        section = doc.sections[0]
        section.header.paragraphs[0].text = "Header đang hiển thị"
        section.first_page_header.paragraphs[0].text = "Header ẩn cũ"
        section.footer.paragraphs[0].text = "Footer đang hiển thị"
        section.first_page_footer.paragraphs[0].text = "Footer ẩn cũ"
        self.assertFalse(section.different_first_page_header_footer)
        e = Editor(save(doc), {"page_numbers": "arabic"}).run()
        result = Document(BytesIO(e.package.serialize())).sections[0]
        self.assertEqual(
            result.first_page_header.paragraphs[0].text, "Header đang hiển thị"
        )
        self.assertIn(
            "Footer đang hiển thị",
            "".join(p.text for p in result.first_page_footer.paragraphs),
        )
        self.assertNotIn(
            "Footer ẩn cũ", "".join(p.text for p in result.first_page_footer.paragraphs)
        )

    def test_existing_page_field_split_across_runs_is_not_duplicated(self):
        doc = Document()
        doc.add_paragraph("Nội dung")
        p = doc.sections[0].footer.paragraphs[0]
        p.add_run("Trang ")
        for kind, text in [
            ("fldChar", "begin"),
            ("instrText", " PA"),
            ("instrText", "GE "),
            ("fldChar", "separate"),
            ("t", "1"),
            ("fldChar", "end"),
        ]:
            run = OxmlElement("w:r")
            node = OxmlElement("w:" + kind)
            if kind == "fldChar":
                node.set(q("fldCharType"), text)
            else:
                node.text = text
            run.append(node)
            p._p.append(run)
        e = Editor(save(doc), {"page_numbers": "arabic"}).run()
        out = Document(BytesIO(e.package.serialize())).sections[0]
        self.assertEqual(len(out.footer._element.xpath(".//w:fldSimple")), 0)
        self.assertEqual(
            len(out.footer._element.xpath(".//w:fldChar[@w:fldCharType='begin']")), 1
        )
        self.assertEqual(len(out.first_page_footer._element.xpath(".//w:fldChar")), 0)
        self.assertEqual(out.first_page_footer.paragraphs[0].text, "Trang ")

    def test_preview_missing_is_partial_not_success(self):
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch(
                "v1.pipeline.render_pdf",
                side_effect=PreviewError("Cần cài LibreOffice"),
            ),
        ):
            out, report = format_document(fixture(1), directory=tmp)
            self.assertEqual(report["status"], "partial")
            self.assertFalse(report["preview"]["available"])
            self.assertTrue(report["integrity"]["verified"])
            Document(BytesIO(out))

    @unittest.skipUnless(find_office(), "LibreOffice is required for integration test")
    def test_real_pdf_toc_pages_and_png(self):
        data = fixture(3)
        with tempfile.TemporaryDirectory() as tmp:
            out, report = format_document(
                data,
                {
                    "mode": "layout",
                    "heading_levels": approved(data),
                    "toc": "insert",
                    "page_numbers": "arabic",
                },
                tmp,
            )
            self.assertTrue(report["preview"]["available"])
            self.assertEqual(report["toc_status"], "verified")
            self.assertGreaterEqual(report["preview"]["page_count"], 2)
            image = Path(tmp) / "page.png"
            page_png(Path(tmp) / "result.pdf", 1, image)
            self.assertTrue(image.read_bytes().startswith(b"\x89PNG"))
            Document(BytesIO(out))

    def test_renderer_rejects_external_loading_fields(self):
        d = Document()
        p = d.add_paragraph()
        f = OxmlElement("w:fldSimple")
        f.set(q("instr"), ' INCLUDETEXT "file:///etc/passwd" ')
        p._p.append(f)
        with self.assertRaises(PreviewError):
            _safe_preview_copy(save(d))

    def test_reviewed_leading_tab_preserves_internal_tabs(self):
        d = Document()
        d.add_paragraph("\tKhối  nội dung\tvẫn nguyên.")
        data = save(d)
        with self.assertRaises(DocumentError):
            Editor(data, {"layout_decisions": {"p00001": "center"}})
        e = Editor(
            data, {"mode": "layout", "layout_decisions": {"p00001": "center"}}
        ).run()
        self.assertEqual(len(e.package.document.xpath("//w:tab", namespaces=NS)), 1)
        self.assertEqual(
            e.package.document.xpath("//w:t/text()", namespaces=NS),
            Package(data).document.xpath("//w:t/text()", namespaces=NS),
        )
        self.assertEqual(e.package.document.find(".//w:jc", NS).get(q("val")), "center")
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_short_manual_alignment_requires_explicit_decision(self):
        d = Document()
        d.add_paragraph("                  Một khối ngắn")
        data = save(d)
        e = Editor(data, {"mode": "layout"}).run()
        self.assertEqual(
            e.package.document.xpath("//w:t/text()", namespaces=NS),
            Package(data).document.xpath("//w:t/text()", namespaces=NS),
        )
        e = Editor(
            data, {"mode": "layout", "layout_decisions": {"p00001": "right_block"}}
        ).run()
        self.assertEqual(
            e.package.document.xpath("//w:t/text()", namespaces=NS), ["Một khối ngắn"]
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_portrait_image_is_limited_by_height_and_warns_distortion(self):
        from v1_fixtures import PNG
        from docx.shared import Inches

        d = Document()
        d.add_picture(BytesIO(PNG), width=Inches(2), height=Inches(16))
        data = save(d)
        e = Editor(data).run()
        extent = e.package.document.xpath(
            "//wp:extent",
            namespaces={
                **NS,
                "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
            },
        )[0]
        self.assertLessEqual(
            int(extent.get("cy")),
            (e._available_height(e.analysis.blocks[0]) - 480) * 635,
        )
        self.assertIn("image_aspect", [w["code"] for w in e.warnings])
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_renderer_timeout_stops_own_process(self):
        import os, sys, time
        from v1.renderer import render_pdf

        if os.name == "nt":
            self.skipTest("POSIX fake executable used for cancellation test")
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            exe = root / "slow-office"
            pidfile = root / "pid"
            exe.write_text(
                "#!"
                + sys.executable
                + "\nimport os,time\nfrom pathlib import Path\nPath("
                + repr(str(pidfile))
                + ").write_text(str(os.getpid()))\ntime.sleep(30)\n"
            )
            exe.chmod(0o700)
            started = time.monotonic()
            with (
                patch("v1.renderer.find_office", return_value=str(exe)),
                self.assertRaises(PreviewError),
            ):
                render_pdf(fixture(1), root / "test.pdf", timeout=0.1)
            self.assertLess(time.monotonic() - started, 3)
            if pidfile.is_file():
                with self.assertRaises(ProcessLookupError):
                    os.kill(int(pidfile.read_text()), 0)

    def test_preview_does_not_silently_drop_external_images(self):
        from v1.package import PKG, R

        package = Package(fixture(1))
        name = "word/_rels/document.xml.rels"
        root = package.xml(name)
        package.write_roots.add(name)
        rel = etree.SubElement(
            root,
            "{" + PKG + "}Relationship",
            {
                "Id": "rIdExternalTest",
                "Type": R + "/hyperlink",
                "TargetMode": "External",
                "Target": "https://example.com/private",
            },
        )
        safe = _safe_preview_copy(package.serialize())
        with zipfile.ZipFile(BytesIO(safe)) as archive:
            self.assertIn(b"about:blank", archive.read(name))
        self.assertEqual(rel.get("Target"), "https://example.com/private")
        rel.set("Type", R + "/image")
        with self.assertRaises(PreviewError):
            _safe_preview_copy(package.serialize())

    def test_svg_guard_parses_utf16_and_allows_basic_vector(self):
        from v1.renderer import validate_svg

        validate_svg(
            b'<svg xmlns="http://www.w3.org/2000/svg"><rect width="10" height="10" fill="green"/></svg>'
        )
        malicious = '<svg xmlns="http://www.w3.org/2000/svg"><image href="file:///etc/passwd"/></svg>'.encode(
            "utf-16"
        )
        with self.assertRaises(PreviewError):
            validate_svg(malicious)
