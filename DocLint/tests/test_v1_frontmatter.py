from contextlib import closing
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
import pypdfium2 as pdfium
from core.xmlutils import NS, q
from v1.analysis import Analysis, plain
from v1.editor import Editor
from v1.package import Package, canonical
from v1.pipeline import format_document
from v1.renderer import (
    find_office,
    render_pdf,
    inspect_pdf,
    PreviewError,
    _office_workspace,
    cleanup_office_workspaces,
)
from v1_fixtures import save

W14 = "http://schemas.microsoft.com/office/word/2010/wordml"


def sample(explicit_break=False):
    d = Document()
    d.add_heading("BÁO CÁO COVER_SENTINEL", 1)
    for index in range(8):
        p = d.add_paragraph()._p
        p.set(f"{{{W14}}}paraId", f"A00000{index:02X}")
        p.set(f"{{{W14}}}textId", "77777777")
        p.set(q("rsidR"), "12345678")
    if explicit_break:
        d.add_page_break()
    d.add_paragraph(
        "INTRODUCTION_SENTINEL belongs to the body, before its first heading."
    )
    d.add_heading("CHƯƠNG 1. SECTION_SENTINEL", 1)
    d.add_paragraph("BODY_SENTINEL must stay on the content page.")
    return save(d)


def page_texts(path):
    result = []
    with closing(pdfium.PdfDocument(path)) as pdf:
        for page in pdf:
            with closing(page), closing(page.get_textpage()) as text:
                result.append(plain(text.get_text_bounded()))
    return result


class FrontmatterTests(unittest.TestCase):
    @unittest.skipUnless(find_office(), "LibreOffice is required")
    def test_cover_toc_and_body_have_distinct_pages(self):
        data = sample()
        a = Analysis(Package(data))
        with tempfile.TemporaryDirectory() as tmp:
            output, report = format_document(
                data,
                {
                    "cover": "academic",
                    "cover_end": a.cover_end,
                    "mode": "layout",
                    "toc": "insert",
                    "page_numbers": "arabic",
                },
                tmp,
            )
            self.assertEqual(report["status"], "ready")
            self.assertEqual(report["toc_status"], "verified")
            self.assertTrue(report["integrity"]["verified"])
            pages = page_texts(Path(tmp) / "result.pdf")
            self.assertEqual(len(pages), 3)
            self.assertIn("COVER_SENTINEL", pages[0])
            self.assertNotIn("MỤC LỤC", pages[0])
            self.assertIn("MỤC LỤC", pages[1])
            self.assertNotIn("COVER_SENTINEL", pages[1])
            self.assertNotIn("INTRODUCTION_SENTINEL", pages[1])
            self.assertIn("INTRODUCTION_SENTINEL", pages[2])
            self.assertIn("BODY_SENTINEL", pages[2])
            # The download must retain the same cached number as the PDF, and
            # must not ask Word to discard that cache while opening the file.
            exported = Package(output)
            toc = exported.document.xpath(
                './/w:sdt[w:sdtPr/w:tag[@w:val="DocLintTOC"]]', namespaces=NS
            )[0]
            self.assertEqual(toc.xpath(".//w:t/text()", namespaces=NS)[-1], "2")
            self.assertIn("2", pages[1])
            self.assertEqual(
                toc.xpath(
                    './/w:fldChar[@w:fldCharType="begin"]/@w:dirty', namespaces=NS
                ),
                ["false"],
            )
            self.assertFalse(toc.xpath(".//@w:fldLock", namespaces=NS))
            self.assertEqual(
                exported.xml("word/settings.xml").xpath(
                    "./w:updateFields/@w:val", namespaces=NS
                ),
                ["false"],
            )
            self.assertEqual(output, (Path(tmp) / "result.docx").read_bytes())

    def test_unverified_toc_keeps_pending_cache_and_allows_manual_update(self):
        output, report = format_document(sample(), {"toc": "insert"}, preview=False)
        self.assertEqual(report["toc_status"], "unresolved")
        exported = Package(output)
        toc = exported.document.xpath(
            './/w:sdt[w:sdtPr/w:tag[@w:val="DocLintTOC"]]', namespaces=NS
        )[0]
        self.assertEqual(toc.xpath(".//w:t/text()", namespaces=NS)[-1], "…")
        self.assertEqual(
            toc.xpath('.//w:fldChar[@w:fldCharType="begin"]/@w:dirty', namespaces=NS),
            ["true"],
        )
        self.assertFalse(toc.xpath(".//@w:fldLock", namespaces=NS))
        self.assertIn("toc_pages_pending", [w["code"] for w in report["warnings"]])
        self.assertTrue(report["integrity"]["verified"])

    def test_generated_toc_defers_even_an_existing_update_on_open_setting(self):
        source = Package(sample())
        root = source.xml("word/settings.xml")
        node = root.makeelement(q("updateFields"), {q("val"): "true"})
        root.insert(0, node)
        source.write_roots.add("word/settings.xml")
        editor = Editor(source.serialize(), {"toc": "insert"}).run()
        self.assertEqual(
            editor.package.xml("word/settings.xml")
            .find(q("updateFields"))
            .get(q("val")),
            "false",
        )
        self.assertTrue(editor.ledger.verify_bytes(editor.package.serialize()))

    def test_page_numbers_do_not_force_original_fields_to_refresh_on_open(self):
        for update_on_open in (None, "false", "true"):
            with self.subTest(update_on_open=update_on_open):
                source = Package(sample())
                root = source.xml("word/settings.xml")
                if update_on_open is not None:
                    root.insert(
                        0,
                        root.makeelement(q("updateFields"), {q("val"): update_on_open}),
                    )
                    source.write_roots.add("word/settings.xml")
                data = source.serialize()
                editor = Editor(data, {"page_numbers": "arabic"}).run()
                output = Package(editor.package.serialize())
                self.assertEqual(
                    canonical(output.xml("word/settings.xml")),
                    canonical(Package(data).xml("word/settings.xml")),
                )
                self.assertTrue(editor.numbering_applied)
                self.assertTrue(editor.ledger.verify_bytes(output.serialize()))

    @unittest.skipUnless(find_office(), "LibreOffice is required")
    def test_toc_precedes_intro_without_a_cover(self):
        d = Document()
        d.add_paragraph("INTRODUCTION_SENTINEL before the first heading.")
        d.add_heading("SECTION_SENTINEL", 1)
        d.add_paragraph("BODY_SENTINEL")
        with tempfile.TemporaryDirectory() as tmp:
            _, report = format_document(save(d), {"toc": "insert"}, tmp)
            self.assertTrue(report["integrity"]["verified"])
            pages = page_texts(Path(tmp) / "result.pdf")
            self.assertEqual(len(pages), 2)
            self.assertIn("MỤC LỤC", pages[0])
            self.assertNotIn("INTRODUCTION_SENTINEL", pages[0])
            self.assertIn("INTRODUCTION_SENTINEL", pages[1])

    @unittest.skipUnless(find_office(), "LibreOffice is required")
    def test_existing_cover_page_break_does_not_add_a_blank_page(self):
        data = sample(explicit_break=True)
        a = Analysis(Package(data))
        with tempfile.TemporaryDirectory() as tmp:
            _, report = format_document(
                data,
                {
                    "cover": "minimal",
                    "cover_end": a.cover_end,
                    "mode": "layout",
                    "toc": "insert",
                },
                tmp,
            )
            pages = page_texts(Path(tmp) / "result.pdf")
            self.assertEqual(len(pages), 3)
            self.assertTrue(report["integrity"]["verified"])

    def test_toc_scope_excludes_an_unchanged_cover_heading(self):
        data = sample()
        e = Editor(data, {"toc": "insert", "mode": "layout"}).run()
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))
        self.assertEqual(
            canonical(e.analysis.elements["p00001"]),
            canonical(Analysis(Package(data)).elements["p00001"]),
        )
        instructions = e.generated_toc.xpath(".//w:instrText/text()", namespaces=NS)
        self.assertIn('\\b "DocLintBody', instructions[0])
        self.assertNotIn(
            "COVER_SENTINEL",
            " ".join(e.generated_toc.xpath(".//w:t/text()", namespaces=NS)),
        )
        self.assertEqual(
            len(e.package.document.xpath(".//w:bookmarkStart", namespaces=NS)), 1
        )

    @unittest.skipUnless(find_office(), "LibreOffice is required")
    def test_locked_intro_starts_after_toc_without_modifying_its_properties(self):
        d = Document()
        d.add_paragraph("INTRODUCTION_SENTINEL must stay exactly unchanged.")
        d.add_heading("SECTION_SENTINEL", 1)
        data = save(d)
        e = Editor(data, {"toc": "insert", "locked_blocks": ["p00001"]}).run()
        self.assertEqual(
            canonical(e.analysis.elements["p00001"]),
            canonical(Analysis(Package(data)).elements["p00001"]),
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))
        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "locked.pdf"
            render_pdf(e.package.serialize(), pdf)
            pages = page_texts(pdf)
            self.assertEqual(len(pages), 2)
            self.assertIn("INTRODUCTION_SENTINEL", pages[1])

    def test_word_editing_metadata_does_not_prevent_removing_unreferenced_padding(self):
        data = sample()
        e = Editor(data, {"cover": "academic", "mode": "layout"}).run()
        original = Analysis(Package(data))
        self.assertEqual(len(original.blocks), 12)
        self.assertEqual(len(e.package.document.findall(".//" + q("p"))), 4)
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_external_paragraph_id_reference_preserves_the_padding_node(self):
        p = Package(sample())
        p.parts["customXml/item1.xml"] = b'<data paraId="A0000000" />'
        e = Editor(p.serialize(), {"cover": "academic", "mode": "layout"}).run()
        self.assertEqual(
            len(
                e.package.document.xpath(
                    ".//w:p[@w14:paraId='A0000000']", namespaces={**NS, "w14": W14}
                )
            ),
            1,
        )
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_cover_boundary_remains_valid_after_chapter_padding_is_removed(self):
        d = Document()
        d.add_heading("BÁO CÁO MÔN HỌC", 1)
        for _ in range(3):
            d.add_paragraph()
        d.add_heading("CHƯƠNG 1. NỘI DUNG", 1)
        d.add_paragraph("Thân bài vẫn giữ nguyên.")
        data = save(d)
        self.assertEqual(Analysis(Package(data)).cover_end, "p00001")
        e = Editor(
            data,
            {
                "cover": "academic",
                "cover_end": "p00004",
                "mode": "layout",
                "toc": "insert",
            },
        ).run()
        self.assertTrue(e.ledger.verify_bytes(e.package.serialize()))

    def test_cleanup_error_does_not_mask_render_error_and_is_retried(self):
        with tempfile.TemporaryDirectory() as tmp:
            marker = PreviewError("Controlled render failure")
            with (
                patch(
                    "v1.renderer.shutil.rmtree",
                    side_effect=OSError(145, "private-path"),
                ),
                self.assertLogs("doclint", "ERROR"),
            ):
                with self.assertRaises(PreviewError) as caught:
                    with _office_workspace(tmp):
                        raise marker
            self.assertIs(caught.exception, marker)
            self.assertTrue(list(Path(tmp).glob("office-*")))
            cleanup_office_workspaces(tmp)
            self.assertFalse(list(Path(tmp).glob("office-*")))

    @unittest.skipUnless(find_office(), "LibreOffice is required")
    def test_cleanup_error_cannot_invalidate_a_successful_pdf(self):
        import shutil

        actual = shutil.rmtree

        def cleanup(path, *args, **kwargs):
            if Path(path).name.startswith("office-"):
                raise OSError(145, "private-path")
            return actual(path, *args, **kwargs)

        with tempfile.TemporaryDirectory() as tmp:
            pdf = Path(tmp) / "result.pdf"
            with (
                patch("v1.renderer.shutil.rmtree", side_effect=cleanup),
                self.assertLogs("doclint", "ERROR"),
            ):
                render_pdf(sample(), pdf)
            self.assertGreater(inspect_pdf(pdf)["page_count"], 0)
            self.assertTrue(list(Path(tmp).glob("office-*")))
            cleanup_office_workspaces(tmp)
            self.assertFalse(list(Path(tmp).glob("office-*")))
