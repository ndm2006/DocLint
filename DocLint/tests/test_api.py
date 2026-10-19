import asyncio
from io import BytesIO
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch
from docx import Document
from fastapi.testclient import TestClient
from api import create_app
from core.processor import MAX_UPLOAD
from core.sessions import SessionStore

class APITests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.app = create_app(self.tmp.name)
        self.client_context = TestClient(self.app); self.client = self.client_context.__enter__()
        d = Document(); d.add_paragraph('Tài liệu thử nghiệm có tiếng Việt, không sửa nội dung.')
        out = BytesIO(); d.save(out); self.data = out.getvalue()

    def tearDown(self):
        self.client_context.__exit__(None, None, None); self.tmp.cleanup()

    def upload(self, name='báo cáo.docx', profile='hust', data=None):
        return self.client.post('/api/v1/lint', files={'file': (name, self.data if data is None else data)}, data={'profile': profile})

    def test_complete_upload_download_delete_flow(self):
        response = self.upload(); self.assertEqual(response.status_code, 201, response.text)
        report = response.json(); self.assertTrue(report['integrity']['verified'])
        self.assertEqual(report['filename'], 'báo cáo_doclint.docx'); self.assertNotIn('expires', report)
        folder = Path(self.tmp.name) / report['session_id']
        self.assertEqual(sorted(p.name for p in folder.iterdir()), ['report.json', 'result.docx'])
        download = self.client.get(report['download_url']); self.assertEqual(download.status_code, 200)
        self.assertEqual(download.headers['cache-control'], 'no-store')
        self.assertIn("filename*=utf-8''", download.headers['content-disposition'])
        self.assertEqual(Document(BytesIO(download.content)).paragraphs[0].text, 'Tài liệu thử nghiệm có tiếng Việt, không sửa nội dung.')
        self.assertEqual(self.client.delete('/api/v1/sessions/' + report['session_id']).status_code, 204)
        self.assertFalse(folder.exists()); self.assertEqual(self.client.get(report['download_url']).status_code, 404)

    def test_validation_and_client_error_messages(self):
        for name, profile, data, expected in [('bad.doc','hust',b'fake',415), ('bad.docx','hust',b'fake',422),('ok.docx','unknown',self.data,422),('empty.docx','hust',b'',422),('big.docx','hust',b'x'*(MAX_UPLOAD+1),413)]:
            response = self.upload(name, profile, data)
            self.assertEqual(response.status_code, expected, response.text)
            self.assertIsInstance(response.json()['detail'], str)
        response = self.client.post('/api/v1/lint', content=b'x', headers={'content-length': str(MAX_UPLOAD + 2 * 1024 * 1024)})
        self.assertEqual(response.status_code, 413, response.text)
        self.assertFalse(list(Path(self.tmp.name).iterdir()))

    def test_chunked_body_without_content_length_is_bounded(self):
        def chunks():
            for _ in range(18):
                yield b'x' * 1024 * 1024
        response = self.client.post('/api/v1/lint', content=chunks(), headers={'content-type': 'application/octet-stream'})
        self.assertEqual(response.status_code, 413, response.text)
        self.assertFalse(list(Path(self.tmp.name).iterdir()))

    def test_profiles_static_security_and_path_rejection(self):
        self.assertEqual({p['id'] for p in self.client.get('/api/v1/profiles').json()}, {'hust','report'})
        for path in ['/', '/static/app.js', '/static/styles.css']:
            response = self.client.get(path); self.assertEqual(response.status_code, 200)
            self.assertIn("script-src 'self'", response.headers['content-security-policy'])
            self.assertEqual(response.headers['x-content-type-options'], 'nosniff')
        self.assertEqual(self.client.get('/api/v1/sessions/not-a-uuid/download').status_code, 404)
        self.assertEqual(self.client.get('/api/v1/health').json(), {'status':'ok'})

    def test_expiry_blocks_download_and_deletes_on_schedule(self):
        with tempfile.TemporaryDirectory() as tmp, TestClient(create_app(tmp, ttl=.1)) as client:
            report = client.post('/api/v1/lint', files={'file': ('x.docx',self.data)}).json()
            client.portal.call(asyncio.sleep, .2)
            self.assertFalse((Path(tmp)/report['session_id']).exists())
            self.assertEqual(client.get(report['download_url']).status_code, 404)

    def test_sessions_are_isolated_and_restart_cleanup_works(self):
        first, second = self.upload().json(), self.upload(profile='report').json()
        self.assertNotEqual(first['session_id'], second['session_id'])
        self.client.delete('/api/v1/sessions/' + first['session_id'])
        self.assertEqual(self.client.get(second['download_url']).status_code, 200)
        with patch('core.sessions.time.time', return_value=time.time()+1900): SessionStore(self.tmp.name).cleanup()
        self.assertFalse(list(Path(self.tmp.name).iterdir()))
