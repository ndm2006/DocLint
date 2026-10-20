from io import BytesIO
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient
from api import create_app
from v1.renderer import find_office, inspect_pdf, PreviewError, _office_environment
from test_v1_jobs import wait_for
from v1_fixtures import fixture


class WindowsOfficeTests(unittest.TestCase):
    def test_windows_prefers_console_for_configured_path_and_standard_install(self):
        with tempfile.TemporaryDirectory() as tmp:
            program = Path(tmp) / "LibreOffice" / "program"
            program.mkdir(parents=True)
            executable = program / "soffice.exe"
            console = program / "soffice.com"
            executable.touch()
            console.touch()
            with (
                patch("v1.renderer.sys.platform", "win32"),
                patch.dict(os.environ, {"DOCLINT_SOFFICE": str(executable)}),
            ):
                self.assertEqual(find_office(), str(console))
                console.unlink()
                self.assertEqual(find_office(), str(executable))
                console.touch()
            with (
                patch("v1.renderer.sys.platform", "win32"),
                patch.dict(os.environ, {"PROGRAMFILES": tmp}, clear=True),
                patch("v1.renderer.shutil.which", return_value=None),
            ):
                self.assertEqual(find_office(), str(console))

    def test_windows_renderer_has_private_appdata_without_inheriting_secrets(self):
        with (
            tempfile.TemporaryDirectory() as tmp,
            patch("v1.renderer.sys.platform", "win32"),
            patch.dict(
                os.environ,
                {
                    "SYSTEMROOT": "C:\\Windows",
                    "PRIVATE_TOKEN": "do-not-inherit",
                    "HTTPS_PROXY": "do-not-inherit",
                },
                clear=True,
            ),
        ):
            root = Path(tmp)
            env = _office_environment(root)
            self.assertEqual(env["SYSTEMROOT"], "C:\\Windows")
            self.assertNotIn("PRIVATE_TOKEN", env)
            self.assertNotIn("HTTPS_PROXY", env)
            for key in ("HOME", "USERPROFILE", "APPDATA", "LOCALAPPDATA"):
                self.assertTrue(Path(env[key]).is_dir())
                self.assertTrue(Path(env[key]).is_relative_to(root))

    def test_pdf_inspection_failure_is_a_preview_error_with_safe_reference(self):
        marker = "private-document-text-and-path"
        with (
            patch("v1.renderer.pdfium.PdfDocument", side_effect=RuntimeError(marker)),
            self.assertLogs("doclint", level="ERROR") as logs,
        ):
            with self.assertRaises(PreviewError) as caught:
                inspect_pdf("not-used.pdf")
        diagnostic = caught.exception.diagnostic
        self.assertEqual(diagnostic["phase"], "inspect_pdf")
        self.assertEqual(diagnostic["exception"], "RuntimeError")
        self.assertIn(diagnostic["reference"], str(caught.exception))
        self.assertNotIn(
            marker, str(caught.exception) + json.dumps(diagnostic) + str(logs.output)
        )


class SupportAPITests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.app = create_app(self.tmp.name)
        self.context = TestClient(self.app)
        self.client = self.context.__enter__()

    def tearDown(self):
        self.context.__exit__(None, None, None)
        self.tmp.cleanup()

    def upload(self):
        source = fixture(1)
        response = self.client.post(
            "/api/v1/jobs", files={"file": ("secret-filename.docx", source)}
        )
        self.assertEqual(response.status_code, 202)
        return source, response.json()["urls"]["status"]

    def test_preview_os_error_does_not_reject_valid_docx_or_block_verified_download(
        self,
    ):
        marker = "private-document-text-and-path"
        with (
            patch(
                "v1.renderer._safe_preview_copy", side_effect=PermissionError(marker)
            ),
            self.assertLogs("doclint", level="ERROR") as logs,
        ):
            source, url = self.upload()
            meta = wait_for(self.client, url)
            self.assertEqual(meta["status"], "review", meta.get("error"))
            self.assertTrue(meta["analysis"])
            self.assertFalse(meta["original_preview"]["available"])
            self.assertIsNone(meta.get("error"))
            diagnostic = self.client.get(url + "/diagnostic")
            self.assertEqual(diagnostic.status_code, 200)
            self.assertEqual(diagnostic.json()["phase"], "render_pdf")
            self.assertEqual(diagnostic.json()["exception"], "PermissionError")
            self.assertIn("attachment", diagnostic.headers["content-disposition"])
            self.assertIn("no-store", diagnostic.headers["cache-control"])
            self.assertNotIn("secret-filename", diagnostic.text)
            self.assertNotIn(marker, diagnostic.text)
            self.assertEqual(self.client.get(url + "/download").status_code, 409)
            self.assertEqual(
                self.client.post(url + "/format", json={}).status_code, 202
            )
            result = wait_for(self.client, url, ("ready", "partial", "review"))
            self.assertEqual(result["status"], "partial", result.get("error"))
            self.assertTrue(result["report"]["integrity"]["verified"])
            self.assertFalse(result["report"]["preview"]["available"])
            self.assertTrue(result["diagnostic"])
            word = self.client.get(url + "/download")
            self.assertEqual(word.status_code, 200)
            self.assertEqual(
                [p.text for p in Document(BytesIO(word.content)).paragraphs],
                [p.text for p in Document(BytesIO(source)).paragraphs],
            )
            self.assertEqual(self.client.delete(url).status_code, 204)
            self.assertEqual(self.client.get(url + "/diagnostic").status_code, 404)
        self.assertNotIn(marker, str(logs.output))

    def test_unexpected_analysis_error_has_diagnostic_and_no_download(self):
        marker = "private-document-text-and-path"
        with (
            patch("v1.jobs.Analysis", side_effect=RuntimeError(marker)),
            self.assertLogs("doclint", level="ERROR") as logs,
        ):
            _, url = self.upload()
            meta = wait_for(self.client, url)
            self.assertEqual(meta["status"], "failed")
            self.assertEqual(meta["diagnostic"]["phase"], "analyze_structure")
            self.assertIn(meta["diagnostic"]["reference"], meta["error"])
            self.assertNotIn(marker, meta["error"])
            detail = self.client.get(url + "/diagnostic")
            self.assertEqual(detail.status_code, 200)
            self.assertNotIn(marker, detail.text)
            self.assertNotIn("secret-filename", detail.text)
            self.assertEqual(self.client.get(url + "/download").status_code, 409)
        self.assertNotIn(marker, str(logs.output))
