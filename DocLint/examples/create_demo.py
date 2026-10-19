"""Generate a synthetic, deliberately messy report; no real user data."""

from io import BytesIO
from pathlib import Path
import argparse
from docx import Document
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw


def demo_bytes():
    doc = Document()
    doc.add_paragraph("TRƯỜNG ĐẠI HỌC MINH HỌA").alignment = 1
    p = doc.add_paragraph("BÁO CÁO MÔN HỌC")
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(24)
    doc.add_paragraph(
        "Đề tài: Thiết kế một công cụ hoàn thiện báo cáo Word"
    ).alignment = 1
    doc.add_paragraph("Nhóm thực hiện: Nhóm DocLint — dữ liệu minh họa")
    doc.add_paragraph("Mã lớp: DEMO-01")
    doc.add_paragraph(
        "Đây là tài liệu tổng hợp để thử phần mềm, không phải báo cáo của một cá nhân hoặc cơ sở đào tạo có thật."
    )
    doc.add_page_break()
    text = (
        "Một báo cáo dễ đọc cần có hệ thống tiêu đề rõ ràng, khoảng cách hợp lý và cách trình bày nhất quán. "
        "Các thao tác căn bằng phím Space, Tab hoặc Enter có thể khiến bố cục thay đổi khi mở tài liệu trên máy khác. "
        "Trong ví dụ này, nội dung cần được giữ nguyên còn thuộc tính trình bày có thể được chuẩn hóa. "
    )
    for chapter, title in [
        (1, "TỔNG QUAN"),
        (2, "PHƯƠNG PHÁP THỬ NGHIỆM"),
        (3, "KẾT QUẢ VÀ ĐÁNH GIÁ"),
    ]:
        if chapter > 1:
            for _ in range(4):
                doc.add_paragraph()
        p = doc.add_paragraph(f"CHƯƠNG {chapter}. {title}")
        p.runs[0].bold = True
        p.runs[0].font.size = Pt(19 if chapter == 1 else 15)
        for section in range(1, 3):
            p = doc.add_paragraph(
                f"{chapter}.{section}. "
                + ("Mục tiêu và phạm vi" if section == 1 else "Yêu cầu cần kiểm chứng")
            )
            p.runs[0].bold = True
            for index in range(2):
                p = doc.add_paragraph(("    " if index == 0 else "") + text)
                run = p.runs[0]
                run.font.name = ["Arial", "Calibri", "Times New Roman"][
                    (chapter + index) % 3
                ]
                run.font.size = Pt([11, 14, 12][(chapter + index) % 3])
                p.paragraph_format.space_after = Pt([0, 18, 3][chapter - 1])
                p.paragraph_format.line_spacing = [1, 1.8, 1.15][chapter - 1]
        if chapter == 1:
            doc.add_paragraph("1. Chuẩn bị bản thảo có lỗi định dạng")
            doc.add_paragraph("2. Duyệt đề xuất và các vùng cần bảo vệ")
            doc.add_paragraph("3. So sánh trước và sau khi chuẩn hóa")
        if chapter == 2:
            doc.add_paragraph("Bảng 2.1. Dữ liệu giả lập để kiểm tra bảng")
            table = doc.add_table(rows=5, cols=3)
            values = [
                ["Mục kiểm tra", "Mô tả", "Giá trị minh họa"],
                ["A", "Định dạng và cấu trúc", "0012"],
                ["B", "Toàn vẹn văn bản, dấu và ký hiệu", "42"],
                ["C", "Bảng và ảnh", "3,14"],
                ["D", "Tính ổn định khi xử lý lại", "100%"],
            ]
            for row, values_row in zip(table.rows, values):
                for cell, value in zip(row.cells, values_row):
                    cell.text = value
            for cell in table.rows[0].cells:
                cell.paragraphs[0].runs[0].bold = True
            for column in table.columns:
                column.width = Inches(3)
            image = Image.new("RGB", (1000, 420), "#f4f5ef")
            draw = ImageDraw.Draw(image)
            draw.text((45, 25), "DEMO: DOCUMENT COMPONENTS", fill="#183f35")
            for i, (label, width, color) in enumerate(
                [
                    ("TEXT", 650, "#183f35"),
                    ("TABLES", 400, "#789270"),
                    ("IMAGES", 300, "#b58966"),
                ]
            ):
                y = 90 + i * 95
                draw.text((45, y + 12), label, fill="#334632")
                draw.rectangle((180, y, 180 + width, y + 50), fill=color)
            buffer = BytesIO()
            image.save(buffer, format="PNG")
            buffer.seek(0)
            doc.add_picture(buffer, width=Inches(7.5))
            doc.add_paragraph("Hình 2.1. Minh họa thành phần tài liệu")
    doc.add_paragraph("                    Sinh viên thực hiện")
    doc.add_heading("TÀI LIỆU THAM KHẢO", level=1)
    doc.add_paragraph(
        "[1] Tài liệu minh họa nội bộ. Không thay đổi nội dung trích dẫn trong quá trình định dạng."
    )
    doc.add_heading("PHỤ LỤC A", level=1)
    p = doc.add_paragraph("    if document_is_valid:\n        preserve_content()")
    p.runs[0].font.name = "Consolas"
    target = BytesIO()
    doc.save(target)
    return target.getvalue()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(
        description="Tạo báo cáo lỗi bố cục để thử DocLint."
    )
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    if args.output.exists():
        parser.error("Tệp đã tồn tại; hãy chọn tên khác.")
    args.output.write_bytes(demo_bytes())
    print(args.output)
